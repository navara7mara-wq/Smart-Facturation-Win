import ctypes
import json
import logging
import os
import shutil
import subprocess
import sys
import threading
import time
import urllib.request
from logging.handlers import RotatingFileHandler
from pathlib import Path

APP_NAME = "PhoEniX BPU"
ROOT_DIR = Path(__file__).resolve().parent


def _desktop_data_root():
    configured = os.environ.get("PHOENIX_USER_DATA_DIR")
    if configured:
        return Path(configured).expanduser().resolve()
    if getattr(sys, "frozen", False):
        local_app_data = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local"))
        return (local_app_data / "SAPTA" / APP_NAME).resolve()
    return None


def _migrate_legacy_data(data_root):
    database_target = data_root / "data" / "pos_ai.sqlite3"
    if database_target.exists():
        return False
    executable_root = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else ROOT_DIR
    legacy_database = executable_root / "data" / "pos_ai.sqlite3"
    if not legacy_database.exists() or legacy_database.resolve() == database_target.resolve():
        return False
    database_target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(legacy_database, database_target)
    for folder_name in ("uploads", "exports", "backups"):
        source = executable_root / folder_name
        target = data_root / folder_name
        if source.exists() and not target.exists():
            shutil.copytree(source, target)
    return True


def configure_runtime_environment():
    data_root = _desktop_data_root()
    if data_root is None:
        return None
    data_root.mkdir(parents=True, exist_ok=True)
    database_path = data_root / "data" / "pos_ai.sqlite3"
    database_existed = database_path.exists()
    legacy_migrated = _migrate_legacy_data(data_root)
    if not database_existed and not legacy_migrated and not database_path.exists():
        os.environ.setdefault("PHOENIX_FRESH_INSTALL_EMPTY", "1")
    defaults = {
        "PHOENIX_DB_PATH": database_path,
        "PHOENIX_UPLOAD_DIR": data_root / "uploads",
        "PHOENIX_EXPORT_DIR": data_root / "exports",
        "PHOENIX_BACKUP_DIR": data_root / "backups",
        "PHOENIX_LICENSE_PATH": data_root / "data" / "license.json",
        "PHOENIX_LOG_DIR": data_root / "logs",
    }
    for key, path in defaults.items():
        os.environ.setdefault(key, str(path))
    return data_root


USER_DATA_ROOT = configure_runtime_environment()

from app import APP_VERSION, BUNDLED_NODE, NODE_MODULES, create_server  # noqa: E402
from db import db  # noqa: E402


LOG_DIR = Path(os.environ.get("PHOENIX_LOG_DIR", ROOT_DIR / "data" / "logs"))


class AlreadyRunningError(RuntimeError):
    pass


class SingleInstanceLock:
    def __init__(self, name="Local\\SAPTA.PhoEniX_BPU.Desktop"):
        self.name = name
        self.handle = None
        self.kernel32 = None

    def acquire(self):
        if os.name != "nt":
            return
        self.kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        self.kernel32.CreateMutexW.argtypes = [ctypes.c_void_p, ctypes.c_bool, ctypes.c_wchar_p]
        self.kernel32.CreateMutexW.restype = ctypes.c_void_p
        self.handle = self.kernel32.CreateMutexW(None, False, self.name)
        if not self.handle:
            raise OSError(ctypes.get_last_error(), "Impossible de creer le verrou de l'application.")
        if ctypes.get_last_error() == 183:
            self.release()
            raise AlreadyRunningError("PhoEniX BPU est deja ouvert.")

    def release(self):
        if self.handle and self.kernel32:
            self.kernel32.CloseHandle(self.handle)
        self.handle = None


class DesktopServer:
    def __init__(self, host="127.0.0.1", port=0):
        self.server = create_server(host, port)
        self.thread = threading.Thread(
            target=self.server.serve_forever,
            name="phoenix-http-server",
            daemon=True,
        )

    @property
    def url(self):
        host, port = self.server.server_address[:2]
        return f"http://{host}:{port}"

    def start(self):
        self.thread.start()

    def stop(self):
        if self.thread.is_alive():
            self.server.shutdown()
            self.thread.join(timeout=5)
        self.server.server_close()


def configure_logging():
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    handler = RotatingFileHandler(
        LOG_DIR / "desktop.log",
        maxBytes=1_000_000,
        backupCount=3,
        encoding="utf-8",
    )
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s"))
    logger = logging.getLogger("phoenix")
    logger.setLevel(logging.INFO)
    if not logger.handlers:
        logger.addHandler(handler)
    return logging.getLogger("phoenix.desktop")


def run_smoke_test():
    started_at = time.perf_counter()
    fresh_install_expected = os.environ.get("PHOENIX_FRESH_INSTALL_EMPTY") == "1"
    server = DesktopServer()
    server.start()
    result = {
        "application": APP_NAME,
        "version": APP_VERSION,
        "url": server.url,
        "database": "unknown",
        "http": "unknown",
        "excel_runtime": "unknown",
        "pdf_runtime": "unknown",
        "webview_runtime": "unknown",
        "fresh_install_empty": "unknown",
        "default_admin": "unknown",
    }
    try:
        with urllib.request.urlopen(server.url + "/login", timeout=15) as response:
            if response.status != 200:
                raise RuntimeError(f"HTTP login returned {response.status}")
            response.read(512)
        with db() as connection:
            integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
            business_tables = (
                "bpu_items", "bpu_st_items", "bpu_st_versions", "clients",
                "client_directions", "company_branches", "purchase_orders",
                "sites", "invoices", "subcontractors", "design_offices",
                "typologies",
            )
            populated = {
                table: connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
                for table in business_tables
            }
            admin = connection.execute(
                "SELECT username, password_hash, must_change_password FROM users WHERE username='admin'"
            ).fetchone()
        if integrity != "ok":
            raise RuntimeError(f"Database integrity check failed: {integrity}")
        if fresh_install_expected:
            if any(populated.values()):
                raise RuntimeError(f"Fresh installation contains business data: {populated}")
            from services.auth import first_run_required, verify_password

            if not admin or verify_password("admin123", admin["password_hash"]):
                raise RuntimeError("Fresh installation contains a reusable default credential")
            if not admin["must_change_password"] or not first_run_required():
                raise RuntimeError("Fresh installation does not require secure administrator setup")
            result["fresh_install_empty"] = "ok"
            result["default_admin"] = "secure-first-run-required"
        else:
            result["fresh_install_empty"] = "not_applicable"
            result["default_admin"] = "preserved"
        from openpyxl import Workbook, load_workbook
        from scripts.export_invoice_template import build as _excel_build

        if not callable(_excel_build):
            raise RuntimeError("Excel export runtime is unavailable")
        load_webview()
        diagnostic_dir = USER_DATA_ROOT or ROOT_DIR / "output"
        diagnostic_dir.mkdir(parents=True, exist_ok=True)
        excel_output = diagnostic_dir / "excel-runtime-smoke.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "PhoEniX BPU"
        workbook.save(excel_output)
        if load_workbook(excel_output, read_only=True).active["A1"].value != "PhoEniX BPU":
            raise RuntimeError("Excel runtime produced an invalid workbook")
        node_environment = os.environ.copy()
        node_environment["PHOENIX_RENDER_ALLOW_LOGIN"] = "1"
        if NODE_MODULES.exists():
            node_environment["NODE_PATH"] = str(NODE_MODULES)
        pdf_output = diagnostic_dir / "pdf-runtime-smoke.pdf"
        node_result = subprocess.run(
            [
                str(BUNDLED_NODE),
                str(ROOT_DIR / "scripts" / "render_pdf.js"),
                server.url + "/login",
                str(pdf_output),
            ],
            capture_output=True,
            text=True,
            timeout=45,
            env=node_environment,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        if node_result.returncode != 0:
            raise RuntimeError(node_result.stderr or "PDF runtime is unavailable")
        if not pdf_output.exists() or not pdf_output.read_bytes().startswith(b"%PDF"):
            raise RuntimeError("PDF runtime produced an invalid document")
        result["database"] = "ok"
        result["http"] = "ok"
        result["excel_runtime"] = "ok"
        result["pdf_runtime"] = "ok"
        result["webview_runtime"] = "ok"
        result["elapsed_ms"] = round((time.perf_counter() - started_at) * 1000)
        return 0, result
    except Exception as error:
        result["error"] = str(error)
        result["elapsed_ms"] = round((time.perf_counter() - started_at) * 1000)
        return 1, result
    finally:
        server.stop()
        output_dir = USER_DATA_ROOT or ROOT_DIR / "output"
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "desktop-smoke-test.json").write_text(
            json.dumps(result, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )


def show_error(message):
    if os.name == "nt":
        ctypes.windll.user32.MessageBoxW(None, str(message), APP_NAME, 0x10)
    else:
        print(message)


def show_info(message):
    if os.name == "nt":
        ctypes.windll.user32.MessageBoxW(None, str(message), APP_NAME, 0x40)
    else:
        print(message)


def focus_existing_window():
    if os.name != "nt":
        return False
    user32 = ctypes.WinDLL("user32", use_last_error=True)
    user32.FindWindowW.argtypes = [ctypes.c_wchar_p, ctypes.c_wchar_p]
    user32.FindWindowW.restype = ctypes.c_void_p
    user32.ShowWindow.argtypes = [ctypes.c_void_p, ctypes.c_int]
    user32.SetForegroundWindow.argtypes = [ctypes.c_void_p]
    window = user32.FindWindowW(None, f"{APP_NAME} {APP_VERSION}")
    if not window:
        window = user32.FindWindowW(None, APP_NAME)
    if not window:
        return False
    user32.ShowWindow(window, 9)
    user32.SetForegroundWindow(window)
    return True


def load_webview():
    try:
        import webview
    except ImportError as error:
        raise RuntimeError(
            "Le composant desktop n'est pas installe. Lancez setup.cmd puis reessayez."
        ) from error
    return webview


def run_desktop(webview_module=None):
    webview = webview_module or load_webview()
    server = DesktopServer()
    server.start()
    try:
        webview.create_window(
            f"{APP_NAME} {APP_VERSION}",
            server.url,
            width=1440,
            height=900,
            min_size=(1024, 700),
            text_select=True,
        )
        webview.start(gui="edgechromium", debug=False)
    finally:
        server.stop()


def main():
    logger = configure_logging()
    if "--smoke-test" in sys.argv:
        exit_code, result = run_smoke_test()
        if exit_code:
            logger.error("Desktop smoke test failed: %s", result.get("error"))
        else:
            logger.info("Desktop smoke test passed in %sms", result["elapsed_ms"])
        return exit_code
    lock = SingleInstanceLock()
    try:
        lock.acquire()
        logger.info("Starting desktop application version %s", APP_VERSION)
        run_desktop()
        logger.info("Desktop application stopped normally")
        return 0
    except AlreadyRunningError as error:
        logger.info("A second launch was redirected to the existing window")
        if not focus_existing_window():
            show_info(error)
        return 0
    except Exception as error:
        logger.exception("Desktop application failed")
        show_error(f"Impossible de demarrer {APP_NAME}.\n\n{error}")
        return 1
    finally:
        lock.release()


if __name__ == "__main__":
    raise SystemExit(main())
