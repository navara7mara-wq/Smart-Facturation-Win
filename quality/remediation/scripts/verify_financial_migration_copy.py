"""Verify v17 -> v19 financial migration on an isolated QA database copy only."""

from __future__ import annotations

import argparse
from contextlib import closing
from decimal import Decimal
from hashlib import sha256
from http.cookiejar import CookieJar
from http.server import ThreadingHTTPServer
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sqlite3
import sys
import threading
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, Request, build_opener


PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
ALLOWED_TARGET_ROOT = (PROJECT_ROOT / "tmp" / "rc2-financial").resolve()
QA_PASSWORD = "RC2MigrationQaOnly-2026"


def file_hash(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def invoice_snapshot(connection: sqlite3.Connection) -> list[list[object]]:
    return [
        list(row)
        for row in connection.execute(
            """
            SELECT id, invoice_number, invoice_type, purchase_order_id, site_id,
                   invoice_date, total_ht, retenue_garantie, montant_ht_apres_rg,
                   tva, total_ttc, montant_en_lettres, remarque, depos,
                   cancelled_at, cancelled_by, cancellation_reason,
                   created_by, updated_by, deleted_at, deleted_by,
                   numbering_system, typologie_snapshot, typologie_label_snapshot
            FROM invoices ORDER BY id
            """
        )
    ]


def line_snapshot(connection: sqlite3.Connection) -> list[list[object]]:
    return [
        list(row)
        for row in connection.execute(
            """
            SELECT invoice_id, article_number, designation_snapshot, unite_snapshot,
                   pu_ht_snapshot, categorie_snapshot, quantite, montant_ht,
                   source_reference_type, source_st_number, mapping_version
            FROM invoice_lines ORDER BY invoice_id, id
            """
        )
    ]


def request(opener, url: str, data: dict[str, str] | None = None, timeout=60):
    payload = urlencode(data).encode() if data is not None else None
    return opener.open(Request(url, data=payload), timeout=timeout)


def csrf(html: str) -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    if not match:
        raise AssertionError("CSRF token missing")
    return match.group(1)


def pdf_exports(target: Path, invoice_ids: list[int], output_dir: Path) -> list[dict[str, object]]:
    os.environ["PHOENIX_DB_PATH"] = str(target)
    from app import App
    from services.auth import hash_password

    with sqlite3.connect(target) as connection:
        connection.execute(
            """
            UPDATE users
            SET password_hash=?, is_active=1, must_change_password=0,
                failed_attempts=0, locked_until=NULL
            WHERE username='admin'
            """,
            (hash_password(QA_PASSWORD),),
        )
        connection.commit()

    server = ThreadingHTTPServer(("127.0.0.1", 0), App)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base_url = f"http://127.0.0.1:{server.server_port}"
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    results = []
    try:
        login_html = request(opener, base_url + "/login").read().decode("utf-8")
        request(
            opener,
            base_url + "/login",
            {"csrf_token": csrf(login_html), "username": "admin", "password": QA_PASSWORD},
        ).read()
        output_dir.mkdir(parents=True, exist_ok=True)
        for invoice_id in invoice_ids:
            content = request(
                opener,
                base_url + f"/invoices/pdf/facture/{invoice_id}/facture.pdf",
                timeout=90,
            ).read()
            assert content.startswith(b"%PDF") and len(content) > 10_000
            output = output_dir / f"historical-invoice-{invoice_id}.pdf"
            output.write_bytes(content)
            results.append(
                {"invoice_id": invoice_id, "path": str(output), "bytes": len(content), "sha256": file_hash(output)}
            )
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=10)
    return results


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    parser.add_argument("evidence", type=Path)
    args = parser.parse_args()

    source = args.source.resolve()
    target = args.target.resolve()
    evidence = args.evidence.resolve()
    if source == target:
        raise RuntimeError("Source and target must differ")
    if ALLOWED_TARGET_ROOT not in target.parents:
        raise RuntimeError(f"QA target outside allowed root: {ALLOWED_TARGET_ROOT}")
    if target.exists():
        raise RuntimeError(f"Refusing to overwrite existing QA target: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    evidence.parent.mkdir(parents=True, exist_ok=True)

    source_hash_before = file_hash(source)
    source_uri = f"file:{source.as_posix()}?mode=ro"
    with closing(sqlite3.connect(source_uri, uri=True)) as source_connection:
        source_version = source_connection.execute(
            "SELECT MAX(version) FROM schema_migrations"
        ).fetchone()[0]
        assert source_version == 17
        with closing(sqlite3.connect(target)) as target_connection:
            source_connection.backup(target_connection)

    from database import migrations

    with closing(sqlite3.connect(target)) as connection:
        connection.row_factory = sqlite3.Row
        before_invoices = invoice_snapshot(connection)
        before_lines = line_snapshot(connection)
        before_relations = [
            list(row)
            for row in connection.execute(
                "SELECT invoice_id, site_id FROM invoice_sites ORDER BY invoice_id, site_id"
            )
        ]
        invoice_ids = [int(row[0]) for row in before_invoices]
        assert invoice_ids
        applied = migrations.run_migrations(connection, target)
        after_invoices = invoice_snapshot(connection)
        after_lines = line_snapshot(connection)
        after_relations = [
            list(row)
            for row in connection.execute(
                "SELECT invoice_id, site_id FROM invoice_sites ORDER BY invoice_id, site_id"
            )
        ]
        rates = [
            list(row)
            for row in connection.execute(
                "SELECT id, rg_rate, tva_rate FROM invoices ORDER BY id"
            )
        ]
        assert applied == [18, 19]
        assert after_invoices == before_invoices
        assert after_lines == before_lines
        assert after_relations == before_relations
        assert all(Decimal(str(row[1])) == Decimal("5") and Decimal(str(row[2])) == Decimal("19") for row in rates)
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
        connection.executemany(
            """
            INSERT INTO app_settings(key, value, updated_by)
            VALUES(?, ?, 'qa-migration-isolation')
            ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_by=excluded.updated_by
            """,
            (("retention_rate", "0.10"), ("tax_rate", "0.20")),
        )
        connection.commit()

    os.environ["PHOENIX_DB_PATH"] = str(target)
    from scripts.export_invoice_template import build as build_xlsx
    from openpyxl import load_workbook

    output_root = evidence.parent / "historical_migration_documents"
    xlsx_dir = output_root / "xlsx"
    xlsx_dir.mkdir(parents=True, exist_ok=True)
    xlsx_results = []
    for invoice_id in invoice_ids:
        output = xlsx_dir / f"historical-invoice-{invoice_id}.xlsx"
        build_xlsx(invoice_id, output)
        workbook = load_workbook(output, read_only=True, data_only=True)
        values = [cell.value for row in workbook["Facture"].iter_rows() for cell in row]
        workbook.close()
        assert "RETENUE DE GARANTIE 5%" in values
        assert "TVA 19 %" in values
        xlsx_results.append(
            {"invoice_id": invoice_id, "path": str(output), "bytes": output.stat().st_size, "sha256": file_hash(output)}
        )

    representative = [invoice_ids[0], invoice_ids[-1]] if len(invoice_ids) > 1 else invoice_ids
    pdf_results = pdf_exports(target, representative, output_root / "pdf")

    with closing(sqlite3.connect(target)) as reopened:
        reopened.row_factory = sqlite3.Row
        max_version = reopened.execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0]
        reopened_rates = [list(row) for row in reopened.execute("SELECT id, rg_rate, tva_rate FROM invoices ORDER BY id")]
        integrity = reopened.execute("PRAGMA integrity_check").fetchone()[0]
        foreign_keys = len(reopened.execute("PRAGMA foreign_key_check").fetchall())
        duplicate_numbers = reopened.execute(
            """
            SELECT COUNT(*) FROM (
                SELECT invoice_number, substr(invoice_date,1,4), COUNT(*) count_rows
                FROM invoices WHERE invoice_number IS NOT NULL
                GROUP BY invoice_number, substr(invoice_date,1,4) HAVING count_rows > 1
            )
            """
        ).fetchone()[0]
        orphan_lines = reopened.execute(
            "SELECT COUNT(*) FROM invoice_lines l LEFT JOIN invoices i ON i.id=l.invoice_id WHERE i.id IS NULL"
        ).fetchone()[0]

    source_hash_after = file_hash(source)
    assert source_hash_after == source_hash_before
    assert max_version == 19 and integrity == "ok" and foreign_keys == 0
    assert reopened_rates == rates
    assert duplicate_numbers == 0 and orphan_lines == 0

    result = {
        "source": str(source),
        "source_sha256_before": source_hash_before,
        "source_sha256_after": source_hash_after,
        "source_unchanged": source_hash_before == source_hash_after,
        "source_schema_version": source_version,
        "qa_target": str(target),
        "qa_target_sha256": file_hash(target),
        "applied_migrations": applied,
        "target_schema_version": max_version,
        "invoice_count_before": len(before_invoices),
        "invoice_count_after": len(after_invoices),
        "invoice_financial_values_unchanged": after_invoices == before_invoices,
        "invoice_lines_unchanged": after_lines == before_lines,
        "invoice_site_relations_unchanged": after_relations == before_relations,
        "backfill_rates": rates,
        "global_defaults_changed_after_migration": {"rg_rate": "10", "tva_rate": "20"},
        "historical_snapshots_after_global_change": reopened_rates,
        "xlsx_exports": xlsx_results,
        "pdf_exports": pdf_results,
        "integrity_check": integrity,
        "foreign_key_violations": foreign_keys,
        "duplicate_invoice_numbers": duplicate_numbers,
        "orphan_invoice_lines": orphan_lines,
        "restart_reopen": "PASS",
        "result": "PASS",
    }
    evidence.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
