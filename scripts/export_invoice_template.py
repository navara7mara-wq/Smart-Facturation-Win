from copy import copy
from pathlib import Path
import sys
import json
import tempfile

from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.styles import Alignment, Font, PatternFill, Border
from PIL import Image, ImageDraw, ImageFont

ROOT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT_DIR))

import app  # noqa: E402


TEMPLATE_PATH = ROOT_DIR / "templates" / "Facture_Construction.xlsx"
EXPORT_SETTINGS = dict(app.DEFAULT_TEMPLATE_SETTINGS)
VISUAL_BLOCKS = {}


def setting_int(key):
    try:
        return int(float(EXPORT_SETTINGS.get(key, app.DEFAULT_TEMPLATE_SETTINGS[key])))
    except (TypeError, ValueError):
        return int(float(app.DEFAULT_TEMPLATE_SETTINGS[key]))


def setting_float(key):
    try:
        return float(EXPORT_SETTINGS.get(key, app.DEFAULT_TEMPLATE_SETTINGS[key]))
    except (TypeError, ValueError):
        return float(app.DEFAULT_TEMPLATE_SETTINGS[key])


def setting_text(key):
    return str(EXPORT_SETTINGS.get(key, app.DEFAULT_TEMPLATE_SETTINGS.get(key, "")) or "")


def hide_empty_sections():
    return True


def load_visual_blocks(document_type):
    key = f"visual_blocks_{document_type}_json"
    raw = setting_text(key)
    if not raw and document_type == "facture":
        raw = setting_text("visual_blocks_json")
    try:
        blocks = json.loads(raw) if raw else app.default_visual_blocks(document_type)
    except Exception:
        blocks = app.default_visual_blocks(document_type)
    existing_ids = {str(block.get("id")) for block in blocks}
    for default_block in app.default_visual_blocks(document_type):
        if str(default_block.get("id")) not in existing_ids:
            blocks.insert(0, default_block)
    return {str(block.get("id")): block for block in blocks}


def visual_block(document_type, block_id):
    return VISUAL_BLOCKS.get(document_type, {}).get(block_id, {})


def block_content(document_type, block_id, fallback):
    return str(visual_block(document_type, block_id).get("content") or fallback)


def block_number(document_type, block_id, field, fallback):
    try:
        return float(visual_block(document_type, block_id).get(field, fallback))
    except (TypeError, ValueError):
        return float(fallback)


def table_width_scale(document_type):
    width = block_number(document_type, "articles_table", "w", 740)
    return max(0.3, min(1.8, width / 740))


def column_pixels(ws, column):
    letter = ws.cell(1, column).column_letter
    width = ws.column_dimensions[letter].width or 8.43
    return int(width * 7 + 5)


def row_pixels(ws, row):
    height = ws.row_dimensions[row].height or 15
    return int(height * 96 / 72)


def pixel_anchor(ws, x_px, y_px):
    remaining_x = max(0, int(x_px))
    col = 1
    while col < ws.max_column and remaining_x > column_pixels(ws, col):
        remaining_x -= column_pixels(ws, col)
        col += 1
    remaining_y = max(0, int(y_px))
    row = 1
    while row < ws.max_row and remaining_y > row_pixels(ws, row):
        remaining_y -= row_pixels(ws, row)
        row += 1
    return ws.cell(row, col).coordinate, remaining_x, remaining_y


def clone_row(ws, source_row, target_row):
    ws.insert_rows(target_row)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, ws.max_column + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)


def replace_placeholders(ws, mapping):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                if "{{#each" in cell.value or "{{/each}}" in cell.value:
                    cell.value = ""
                    continue
                value = cell.value
                for key, replacement in mapping.items():
                    value = value.replace("{{" + key + "}}", str(replacement or ""))
                cell.value = value


def write_lines(ws, start_row, start_col, text):
    for offset, line in enumerate(str(text).splitlines()):
        cell = ws.cell(start_row + offset, start_col)
        if cell.__class__.__name__ != "MergedCell":
            cell.value = line


def set_merged_text(ws, row, start_col, end_col, text):
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= start_col <= merged_range.max_col:
            ws.unmerge_cells(str(merged_range))
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    ws.cell(row, start_col).value = text


def apply_visual_text_blocks(ws, document_type):
    if document_type == "facture":
        write_lines(ws, 7, 1, block_content(document_type, "enterprise_info", "RGC: {{Entreprise.RGC}}\nNIF: {{Entreprise.NIF}}\nART: {{Entreprise.ART}}\nADRESSE: {{Entreprise.Adresse}}\nN° COMPTE: {{Entreprise.RIB}}"))
        write_lines(ws, 7, 4, block_content(document_type, "client_info", "DOIT : {{Client.Nom}}\n{{Client.Direction}}\n{{Client.Adresse}}\nRGC N°: {{Client.RGC}}\nNIF N°: {{Client.NIF}}"))
        set_merged_text(ws, 13, 1, 6, block_content(document_type, "invoice_title", "FACTURE N° : {{N_Facture}}"))
        write_lines(ws, 15, 1, block_content(document_type, "site_info", "Référence Contrat: {{Ref_Contrat}}\nCode de site: {{Code_Site}}\nNom de site: {{Nom_Site}}\nRégion: {{Region}}\nTypologie de site: {{Typologie}}\nBon de commande: {{N_BC}}"))
        signature = block_content(document_type, "signature", "L'ENTREPRISE/{{Entreprise.Nom}}")
        for row in range(1, ws.max_row + 1):
            if isinstance(ws.cell(row, 1).value, str) and "L'ENTREPRISE" in ws.cell(row, 1).value:
                ws.cell(row, 1).value = signature
                break
        set_table_header_from_block(ws, document_type, "articles_table")
        set_totals_labels_from_block(ws, document_type)
    else:
        write_lines(ws, 7, 1, block_content(document_type, "site_info", "Code de site: {{Code_Site}}\nNom de site: {{Nom_Site}}\nRégion: {{Region}}\nTypologie de site: {{Typologie}}"))
        set_table_header_from_block(ws, document_type, "articles_table")
        if document_type == "devis_estimatif":
            signature = block_content(document_type, "signature", "L'ENTREPRISE/{{Entreprise.Nom}}")
            for row in range(1, ws.max_row + 1):
                if isinstance(ws.cell(row, 1).value, str) and "L'ENTREPRISE" in ws.cell(row, 1).value:
                    ws.cell(row, 1).value = signature
                    break


def set_table_header_from_block(ws, document_type, block_id):
    content = block_content(document_type, block_id, "")
    if "|" not in content:
        return
    labels = [part.strip() for part in content.splitlines()[0].split("|")]
    for row in range(1, ws.max_row + 1):
        row_values = [ws.cell(row, col).value for col in range(1, min(7, ws.max_column + 1))]
        if any(str(value or "").strip() in {"N°", "N"} for value in row_values) and any("Désignation" in str(value or "") or "Designation" in str(value or "") for value in row_values):
            for col, label in enumerate(labels, start=1):
                ws.cell(row, col).value = label
            return


def set_totals_labels_from_block(ws, document_type):
    content = block_content(document_type, "totals_table", "")
    labels = [line.strip() for line in content.splitlines() if line.strip()]
    if not labels:
        return
    matched_rows = []
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 5).value
        if isinstance(label, str) and (
            "TOTAL" in label or "RETENUE" in label or "MONTANT HT" in label or "T V A" in label or "TVA" in label
        ):
            matched_rows.append(row)
    for row, label in zip(matched_rows[-len(labels):], labels):
        ws.cell(row, 5).value = label


def image_anchor(cell_coordinate, width_px, height_px, col_offset_px=0, row_offset_px=0):
    row, col = coordinate_to_tuple(cell_coordinate)
    marker = AnchorMarker(
        col=col - 1,
        row=row - 1,
        colOff=pixels_to_EMU(col_offset_px),
        rowOff=pixels_to_EMU(row_offset_px),
    )
    ext = XDRPositiveSize2D(pixels_to_EMU(width_px), pixels_to_EMU(height_px))
    return OneCellAnchor(_from=marker, ext=ext)


def place_logo(ws, placeholder, image_path, anchor_coordinate=None, col_offset_px=0, row_offset_px=0, width_px=None, height_px=None):
    if not image_path:
        return
    path = Path(image_path)
    if not path.is_absolute():
        path = ROOT_DIR / path
    if not path.exists():
        return
    coordinate = anchor_coordinate
    logo_width = int(width_px or setting_int("logo_width_px"))
    logo_height = int(height_px or setting_int("logo_height_px"))
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "{{" + placeholder + "}}":
                cell.value = ""
                image = ExcelImage(str(path))
                image.width = logo_width
                image.height = logo_height
                coordinate = anchor_coordinate or cell.coordinate
                if col_offset_px or row_offset_px:
                    image.anchor = image_anchor(coordinate, logo_width, logo_height, col_offset_px=col_offset_px, row_offset_px=row_offset_px)
                    ws.add_image(image)
                else:
                    ws.add_image(image, coordinate)
                return
    if coordinate:
        image = ExcelImage(str(path))
        image.width = logo_width
        image.height = logo_height
        if col_offset_px or row_offset_px:
            image.anchor = image_anchor(coordinate, logo_width, logo_height, col_offset_px=col_offset_px, row_offset_px=row_offset_px)
            ws.add_image(image)
        else:
            ws.add_image(image, coordinate)


def place_logo_from_block(ws, placeholder, image_path, document_type, block_id, fallback_anchor):
    block = visual_block(document_type, block_id)
    if block:
        coordinate, col_offset, row_offset = pixel_anchor(ws, block.get("x", 0), block.get("y", 0))
        place_logo(
            ws,
            placeholder,
            image_path,
            anchor_coordinate=coordinate,
            col_offset_px=col_offset,
            row_offset_px=row_offset,
            width_px=block.get("w"),
            height_px=block.get("h"),
        )
        return
    place_logo(ws, placeholder, image_path, anchor_coordinate=fallback_anchor)


def block_logo_size(document_type, block_id):
    block = visual_block(document_type, block_id)
    if block:
        return block.get("w", 150), block.get("h", 70)
    return 150, 70


def clear_devis_header(ws):
    ws._images = []
    for row in range(1, 7):
        for col in range(1, 7):
            cell = ws.cell(row, col)
            if cell.__class__.__name__ != "MergedCell":
                cell.value = None
                cell.fill = PatternFill(fill_type=None)
                cell.border = Border()
    for row in range(1, 7):
        ws.row_dimensions[row].height = 22


def unmerge_from_row(ws, min_row):
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row >= min_row:
            ws.unmerge_cells(str(merged_range))


def fill_line_row(ws, row_index, line, with_prices=True):
    ws.cell(row_index, 1).value = line["article_number"]
    ws.cell(row_index, 2).value = line["designation_snapshot"]
    ws.cell(row_index, 2).alignment = copy(ws.cell(row_index, 2).alignment)
    ws.cell(row_index, 2).alignment = Alignment(
        horizontal=ws.cell(row_index, 2).alignment.horizontal,
        vertical="top",
        text_rotation=ws.cell(row_index, 2).alignment.text_rotation,
        wrap_text=True,
        shrink_to_fit=False,
        indent=ws.cell(row_index, 2).alignment.indent,
    )
    ws.cell(row_index, 3).value = line["unite_snapshot"]
    ws.cell(row_index, 4).value = line["quantite"]
    if with_prices:
        ws.cell(row_index, 5).value = line["pu_ht_snapshot"]
        ws.cell(row_index, 6).value = line["montant_ht"]
        ws.cell(row_index, 5).number_format = '#,##0.00'
        ws.cell(row_index, 6).number_format = '#,##0.00'
    designation = str(line["designation_snapshot"] or "")
    wraps = "\n" in designation or len(designation) > setting_int("designation_wrap_chars")
    if wraps:
        ws.row_dimensions[row_index].height = setting_float("line_height_wrapped")
    else:
        ws.row_dimensions[row_index].height = setting_float("line_height_normal")


def fill_section(ws, template_row, lines, with_prices=True):
    if not lines:
        for col in range(1, 7 if with_prices else 5):
            cell = ws.cell(template_row, col)
            if cell.__class__.__name__ != "MergedCell":
                cell.value = ""
        return 0
    for offset, line in enumerate(lines):
        row_index = template_row + offset
        if offset:
            clone_row(ws, template_row, row_index)
        fill_line_row(ws, row_index, line, with_prices=with_prices)
    return len(lines)


def grouped_lines(lines):
    return {
        "acquisition": [line for line in lines if line["categorie_snapshot"] == "acquisition"],
        "fourniture": [line for line in lines if line["categorie_snapshot"] == "fourniture"],
        "prestation": [line for line in lines if line["categorie_snapshot"] == "prestation"],
        "ndc": [line for line in lines if line["categorie_snapshot"] == "ndc"],
    }


def build(invoice_id, output_path):
    global EXPORT_SETTINGS, VISUAL_BLOCKS
    EXPORT_SETTINGS = app.template_settings()
    VISUAL_BLOCKS = {
        "facture": load_visual_blocks("facture"),
        "devis_quantitatif": load_visual_blocks("devis_quantitatif"),
        "devis_estimatif": load_visual_blocks("devis_estimatif"),
    }
    invoice, company, contract, lines, ndc_sites = app.invoice_export_data(invoice_id)
    wb = load_workbook(TEMPLATE_PATH)
    groups = grouped_lines(lines)
    site_label = invoice["code_site"] or ", ".join(site["code_site"] for site in ndc_sites)
    site_name = invoice["nom_site"] or f"{len(ndc_sites)} sites NDC"

    mapping = {
        "Entreprise.RGC": company["rgc"],
        "Entreprise.NIF": company["nif"],
        "Entreprise.ART": company["art"],
        "Entreprise.Adresse": company["adresse"],
        "Entreprise.RIB": company["numero_compte"],
        "Entreprise.Nom": company["nom"],
        "Client.Nom": invoice["doit_nom"],
        "Client.Direction": invoice["direction_regionale"],
        "Client.Adresse": invoice["client_adresse"],
        "Client.RGC": invoice["client_rgc"],
        "Client.NIF": invoice["client_nif"],
        "Logo.Client": "",
        "Logo.Entreprise": "",
        "N_Facture": invoice["invoice_number"],
        "FACTURE.TYPE": app.display_type(invoice["invoice_type"]),
        "Ref_Contrat": contract["reference_contrat"],
        "Code_Site": site_label,
        "Nom_Site": site_name,
        "Region": invoice["region"] or "",
        "Typologie": invoice["typologie_site"] or "",
        "N_BC": invoice["numero_bc"],
        "Total_Fournitures_HT": sum(line["montant_ht"] for line in groups["fourniture"] + groups["acquisition"] + groups["ndc"]),
        "Total_Prestations_HT": sum(line["montant_ht"] for line in groups["prestation"]),
        "Montant_Global_HT": invoice["total_ht"],
        "Retenue_Garantie": invoice["retenue_garantie"],
        "Montant_HT_Apres_Retenue": invoice["montant_ht_apres_rg"],
        "Montant_TVA": invoice["tva"],
        "Montant_TTC": invoice["total_ttc"],
        "Montant_En_Lettres": app.amount_to_french(invoice["total_ttc"]).upper(),
    }

    # The official construction template has two material sections: FOURNITURES and PRESTATION.
    # Acquisition/NDC lines are placed in the first section to preserve the official layout.
    first_section = groups["fourniture"] + groups["acquisition"] + groups["ndc"]
    second_section = groups["prestation"]

    fill_facture_sheet(wb["Facture"], first_section, second_section)
    fill_quantitative_sheet(wb["Devis Quantitatif"], first_section, second_section)
    fill_estimative_sheet(wb["Devis Estimatif"], first_section, second_section)
    apply_visual_text_blocks(wb["Facture"], "facture")
    for sheet_name, document_type in (
        ("Devis Quantitatif", "devis_quantitatif"),
        ("Devis Estimatif", "devis_estimatif"),
    ):
        clear_devis_header(wb[sheet_name])
        apply_visual_text_blocks(wb[sheet_name], document_type)
        place_devis_header_image(
            wb[sheet_name],
            document_type,
            mapping,
            invoice["client_logo_path"],
            company["logo_path"],
        )
    place_logo(
        wb["Facture"],
        "Logo.Client",
        invoice["client_logo_path"],
        anchor_coordinate=setting_text("client_logo_anchor") or None,
        col_offset_px=setting_int("client_logo_offset_px"),
    )
    place_logo(
        wb["Facture"],
        "Logo.Entreprise",
        company["logo_path"],
        anchor_coordinate=setting_text("company_logo_anchor") or "E1",
        col_offset_px=setting_int("company_logo_offset_px"),
    )
    for sheet_name in ("Facture", "Devis Quantitatif", "Devis Estimatif"):
        replace_placeholders(wb[sheet_name], mapping)
    format_facture_sheet(wb["Facture"])
    format_devis_sheet(wb["Devis Quantitatif"], with_prices=False)
    format_devis_sheet(wb["Devis Estimatif"], with_prices=True)

    wb.save(output_path)


def fill_facture_sheet(ws, first_section, second_section):
    unmerge_from_row(ws, 21)
    ws["A21"].value = ""
    first_deleted = False
    if first_section:
        first_count = fill_section(ws, 24, first_section, with_prices=True)
        fourniture_total_row = 25 + max(first_count - 1, 0)
        ws.cell(fourniture_total_row, 1).value = ""
        ws.cell(fourniture_total_row, 5).value = "TOTAL  FOURNITURES  HT (01)"
        ws.cell(fourniture_total_row, 6).number_format = '#,##0.00'
    elif hide_empty_sections():
        ws.delete_rows(23, 3)
        first_count = 0
        first_deleted = True
    else:
        first_count = fill_section(ws, 24, first_section, with_prices=True)

    prestation_header_row = 26 + max(first_count - 1, 0)
    if first_deleted:
        prestation_header_row = 23
    prestation_row = prestation_header_row + 1
    if second_section:
        fill_section(ws, prestation_row, second_section, with_prices=True)
        prestation_total_row = prestation_row + 1 + max(len(second_section) - 1, 0)
        ws.cell(prestation_total_row, 1).value = ""
        ws.cell(prestation_total_row, 5).value = "TOTAL  PRESTATION HT (02)"
        ws.cell(prestation_total_row, 6).number_format = '#,##0.00'
    elif hide_empty_sections():
        ws.delete_rows(prestation_header_row, 3)
    else:
        fill_section(ws, prestation_row, second_section, with_prices=True)


def format_facture_sheet(ws):
    scale = table_width_scale("facture")
    widths = {
        "A": setting_float("col_a_width") * scale,
        "B": setting_float("col_b_width") * scale,
        "C": setting_float("col_c_width") * scale,
        "D": setting_float("col_d_width") * scale,
        "E": setting_float("col_e_width") * scale,
        "F": setting_float("col_f_width") * scale,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.sheet_view.view = "normal"
    ws.sheet_view.selection[0].activeCell = "A1"
    ws.sheet_view.selection[0].sqref = "A1"

    # Totals move down when lines are inserted, so scan the visible amount columns.
    for row in range(1, ws.max_row + 1):
        for col in (5, 6):
            cell = ws.cell(row, col)
            coerce_number(cell)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'
        label = ws.cell(row, 5).value
        if isinstance(label, str) and (
            "TOTAL" in label
            or "RETENUE" in label
            or "MONTANT HT" in label
            or "T V A" in label
        ):
            coerce_number(ws.cell(row, 6))
            ws.cell(row, 6).number_format = '#,##0.00'

    for row in range(1, ws.max_row + 1):
        ws.cell(row, 2).alignment = copy(ws.cell(row, 2).alignment)
        ws.cell(row, 2).alignment = Alignment(
            horizontal=ws.cell(row, 2).alignment.horizontal,
            vertical="top",
            wrap_text=True,
            shrink_to_fit=False,
        )

    # The amount-in-words row is immediately below the label in the official template.
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == "Arrêté la présente  Facture  à la somme de:":
            amount_cell = ws.cell(row + 1, 1)
            merge_range = f"A{row + 1}:F{row + 2}"
            for merged_range in list(ws.merged_cells.ranges):
                if merged_range.min_row <= row + 1 <= merged_range.max_row:
                    ws.unmerge_cells(str(merged_range))
            ws.merge_cells(merge_range)
            amount_cell.value = str(amount_cell.value or "").upper()
            amount_cell.font = Font(
                name=amount_cell.font.name,
                size=block_number("facture", "amount_words", "font", setting_int("amount_words_font_size")),
                bold=amount_cell.font.bold,
                italic=amount_cell.font.italic,
                color=amount_cell.font.color,
            )
            amount_cell.alignment = copy(amount_cell.alignment)
            amount_cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )
            ws.row_dimensions[row + 1].height = 34
            ws.row_dimensions[row + 2].height = 34
            break


def replace_text_placeholders(text, mapping):
    value = str(text or "")
    for key, replacement in mapping.items():
        value = value.replace("{{" + key + "}}", str(replacement or ""))
    return value


def make_devis_title_image(document_type, mapping):
    block = visual_block(document_type, "devis_title")
    if not block:
        return None
    width = int(block_number(document_type, "devis_title", "w", 380))
    height = int(block_number(document_type, "devis_title", "h", 72))
    font_size = int(block_number(document_type, "devis_title", "font", 16))
    text = replace_text_placeholders(block_content(document_type, "devis_title", ""), mapping)
    image = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    radius = min(14, height // 4)
    draw.rounded_rectangle((0, 0, width - 1, height - 1), radius=radius, fill=(68, 114, 196, 255), outline=(31, 78, 121, 255), width=2)
    try:
        font = ImageFont.truetype("arialbd.ttf", font_size)
    except OSError:
        try:
            font = ImageFont.truetype("arial.ttf", font_size)
        except OSError:
            font = ImageFont.load_default()
    lines = text.splitlines() or [text]
    line_boxes = [draw.textbbox((0, 0), line, font=font) for line in lines]
    line_heights = [box[3] - box[1] for box in line_boxes]
    total_height = sum(line_heights) + max(0, len(lines) - 1) * 6
    y = max(0, (height - total_height) / 2)
    for line, box, line_height in zip(lines, line_boxes, line_heights):
        text_width = box[2] - box[0]
        draw.text(((width - text_width) / 2, y), line, fill=(255, 255, 255, 255), font=font)
        y += line_height + 6
    target = Path(tempfile.gettempdir()) / f"pos_ai_{document_type}_title.png"
    image.save(target)
    return target, width, height


def place_devis_title_image(ws, document_type, mapping):
    generated = make_devis_title_image(document_type, mapping)
    if not generated:
        return
    path, width, height = generated
    image = ExcelImage(str(path))
    image.width = width
    image.height = height
    # Centered safely between left and right logos.
    image.anchor = image_anchor("B1", width, height, col_offset_px=34, row_offset_px=4)
    ws.add_image(image)


def resolve_image_path(image_path):
    if not image_path:
        return None
    path = Path(image_path)
    if not path.is_absolute():
        path = ROOT_DIR / path
    return path if path.exists() else None


def paste_logo(canvas, image_path, block, default_box):
    path = resolve_image_path(image_path)
    if not path:
        return
    x = int(float(block.get("x", default_box[0]))) if block else default_box[0]
    y = int(float(block.get("y", default_box[1]))) if block else default_box[1]
    width = int(float(block.get("w", default_box[2]))) if block else default_box[2]
    height = int(float(block.get("h", default_box[3]))) if block else default_box[3]
    logo = Image.open(path).convert("RGBA")
    logo.thumbnail((width, height), Image.Resampling.LANCZOS)
    logo_x = x + max(0, (width - logo.width) // 2)
    logo_y = y + max(0, (height - logo.height) // 2)
    canvas.alpha_composite(logo, (logo_x, logo_y))


def draw_title_on_canvas(canvas, document_type, mapping):
    block = visual_block(document_type, "devis_title")
    if not block:
        return
    x = int(block_number(document_type, "devis_title", "x", 205))
    y = int(block_number(document_type, "devis_title", "y", 20))
    width = int(block_number(document_type, "devis_title", "w", 380))
    height = int(block_number(document_type, "devis_title", "h", 72))
    font_size = int(block_number(document_type, "devis_title", "font", 16))
    background = str(block.get("background") or "#4472c4")
    color = str(block.get("color") or "#ffffff")
    align = str(block.get("align") or "center")
    bold = bool(block.get("bold", True))
    text = replace_text_placeholders(block_content(document_type, "devis_title", ""), mapping)
    draw = ImageDraw.Draw(canvas)
    radius = min(14, height // 4)
    def hex_to_rgb(value, fallback):
        value = value.strip().lstrip("#")
        if len(value) == 6:
            try:
                return tuple(int(value[i:i+2], 16) for i in (0, 2, 4)) + (255,)
            except ValueError:
                return fallback
        return fallback
    bg_rgba = hex_to_rgb(background, (68, 114, 196, 255))
    text_rgba = hex_to_rgb(color, (255, 255, 255, 255))
    draw.rounded_rectangle((x, y, x + width - 1, y + height - 1), radius=radius, fill=bg_rgba, outline=(31, 78, 121, 255), width=2)
    try:
        font = ImageFont.truetype("arialbd.ttf" if bold else "arial.ttf", font_size)
    except OSError:
        try:
            font = ImageFont.truetype("arial.ttf", font_size)
        except OSError:
            font = ImageFont.load_default()
    lines = text.splitlines() or [text]
    boxes = [draw.textbbox((0, 0), line, font=font) for line in lines]
    heights = [box[3] - box[1] for box in boxes]
    total_height = sum(heights) + max(0, len(lines) - 1) * 6
    text_y = y + max(0, (height - total_height) / 2)
    for line, box, line_height in zip(lines, boxes, heights):
        text_width = box[2] - box[0]
        if align == "left":
            text_x = x + 12
        elif align == "right":
            text_x = x + width - text_width - 12
        else:
            text_x = x + (width - text_width) / 2
        draw.text((text_x, text_y), line, fill=text_rgba, font=font)
        text_y += line_height + 6


def header_bounds(document_type):
    defaults = {
        "client_logo": (24, 20, 150, 70),
        "devis_title": (205, 20, 380, 72),
        "company_logo": (610, 20, 150, 70),
    }
    boxes = []
    for block_id, default in defaults.items():
        block = visual_block(document_type, block_id)
        x = int(float(block.get("x", default[0]))) if block else default[0]
        y = int(float(block.get("y", default[1]))) if block else default[1]
        w = int(float(block.get("w", default[2]))) if block else default[2]
        h = int(float(block.get("h", default[3]))) if block else default[3]
        boxes.append((x, y, w, h))
    min_x = max(0, min(x for x, _, _, _ in boxes))
    min_y = max(0, min(y for _, y, _, _ in boxes))
    max_x = max(x + w for x, _, w, _ in boxes)
    max_y = max(y + h for _, y, _, h in boxes)
    return min_x, min_y, max_x - min_x, max_y - min_y


def offset_block(block, dx, dy):
    shifted = dict(block or {})
    shifted["x"] = float(shifted.get("x", 0)) - dx
    shifted["y"] = float(shifted.get("y", 0)) - dy
    return shifted


def make_devis_header_image(document_type, mapping, client_logo_path, company_logo_path):
    canvas_width = 794
    canvas_height = 115
    canvas = Image.new("RGBA", (canvas_width, canvas_height), (255, 255, 255, 0))
    paste_logo(canvas, client_logo_path, visual_block(document_type, "client_logo"), (24, 20, 150, 70))
    paste_logo(canvas, company_logo_path, visual_block(document_type, "company_logo"), (610, 20, 150, 70))
    original = VISUAL_BLOCKS[document_type].get("devis_title")
    draw_title_on_canvas(canvas, document_type, mapping)
    if original:
        VISUAL_BLOCKS[document_type]["devis_title"] = original
    target = Path(tempfile.gettempdir()) / f"pos_ai_{document_type}_header.png"
    canvas.save(target)
    return target, canvas_width, canvas_height


def place_devis_header_image(ws, document_type, mapping, client_logo_path, company_logo_path):
    path, width, height = make_devis_header_image(document_type, mapping, client_logo_path, company_logo_path)
    image = ExcelImage(str(path))
    image.width = width
    image.height = height
    ws.add_image(image, "A1")


def coerce_number(cell):
    if isinstance(cell.value, str):
        raw = cell.value.strip().replace(" ", "").replace("\u00a0", "")
        raw = raw.replace(",", ".")
        try:
            number = float(raw)
        except ValueError:
            return
        cell.value = int(number) if number.is_integer() else number


def last_used_row(ws):
    last = 1
    for row in range(1, ws.max_row + 1):
        has_value = any(ws.cell(row, col).value not in (None, "") for col in range(1, min(ws.max_column, 8) + 1))
        if has_value:
            last = row
    return last


def cleanup_devis_sheet(ws, with_prices):
    max_col = 6 if with_prices else 4
    last_row = max(last_used_row(ws) + 3, 42)
    if ws.max_row > last_row:
        ws.delete_rows(last_row + 1, ws.max_row - last_row)
    if ws.max_column > max_col:
        ws.delete_cols(max_col + 1, ws.max_column - max_col)
    ws.print_area = f"A1:{ws.cell(last_row, max_col).coordinate}"
    ws.sheet_view.view = "normal"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.selection[0].activeCell = "A1"
    ws.sheet_view.selection[0].sqref = "A1"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 0
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = 74
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.2
    ws.page_margins.bottom = 0.2


def format_devis_sheet(ws, with_prices):
    document_type = "devis_estimatif" if with_prices else "devis_quantitatif"
    scale = table_width_scale(document_type)
    widths = {
        "A": setting_float("col_a_width") * scale,
        "B": setting_float("col_b_width") * scale,
        "C": setting_float("col_c_width") * scale,
        "D": setting_float("col_d_width") * scale,
    }
    if with_prices:
        widths.update({"E": setting_float("col_e_width") * scale, "F": setting_float("col_f_width") * scale})
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    max_money_col = 6 if with_prices else 4
    for row in range(1, ws.max_row + 1):
        description_cell = ws.cell(row, 2)
        description_cell.alignment = copy(description_cell.alignment)
        description_cell.alignment = Alignment(
            horizontal=description_cell.alignment.horizontal,
            vertical="top",
            wrap_text=True,
            shrink_to_fit=False,
        )
        if with_prices:
            for col in (5, 6):
                cell = ws.cell(row, col)
                coerce_number(cell)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'
            for col in range(1, max_money_col + 1):
                label = ws.cell(row, col).value
                if isinstance(label, str) and "TOTAL" in label:
                    amount_cell = ws.cell(row, 6)
                    coerce_number(amount_cell)
                    amount_cell.number_format = '#,##0.00'

    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == "{{Montant_En_Lettres}}" or (
            isinstance(ws.cell(row, 1).value, str)
            and "DINARS" in ws.cell(row, 1).value
        ):
            amount_cell = ws.cell(row, 1)
            amount_cell.value = str(amount_cell.value or "").upper()
            amount_cell.font = Font(
                name=amount_cell.font.name,
                size=block_number(document_type, "amount_words", "font", setting_int("amount_words_font_size")),
                bold=amount_cell.font.bold,
                italic=amount_cell.font.italic,
                color=amount_cell.font.color,
            )
            amount_cell.alignment = copy(amount_cell.alignment)
            amount_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws.row_dimensions[row].height = 34
            if row + 1 <= ws.max_row:
                ws.row_dimensions[row + 1].height = 34
            break
    cleanup_devis_sheet(ws, with_prices)


def fill_quantitative_sheet(ws, first_section, second_section):
    unmerge_from_row(ws, 12)
    ws["A12"].value = ""
    first_deleted = False
    if first_section:
        first_count = fill_section(ws, 15, first_section, with_prices=False)
    elif hide_empty_sections():
        ws.delete_rows(14, 2)
        first_count = 0
        first_deleted = True
    else:
        first_count = fill_section(ws, 15, first_section, with_prices=False)

    prestation_header_row = 16 + max(first_count - 1, 0)
    if first_deleted:
        prestation_header_row = 14
    prestation_row = prestation_header_row + 1
    if second_section:
        fill_section(ws, prestation_row, second_section, with_prices=False)
    elif hide_empty_sections():
        ws.delete_rows(prestation_header_row, 2)
    else:
        fill_section(ws, prestation_row, second_section, with_prices=False)


def fill_estimative_sheet(ws, first_section, second_section):
    unmerge_from_row(ws, 12)
    ws["A12"].value = ""
    first_deleted = False
    if first_section:
        first_count = fill_section(ws, 15, first_section, with_prices=True)
        fourniture_total_row = 16 + max(first_count - 1, 0)
        ws.cell(fourniture_total_row, 1).value = "TOTAL  FOURNITURES  HT (01)"
        ws.cell(fourniture_total_row, 6).number_format = '#,##0.00'
    elif hide_empty_sections():
        ws.delete_rows(14, 3)
        first_count = 0
        first_deleted = True
    else:
        first_count = fill_section(ws, 15, first_section, with_prices=True)

    prestation_header_row = 17 + max(first_count - 1, 0)
    if first_deleted:
        prestation_header_row = 14
    prestation_row = prestation_header_row + 1
    if second_section:
        fill_section(ws, prestation_row, second_section, with_prices=True)
        prestation_total_row = prestation_row + 1 + max(len(second_section) - 1, 0)
        ws.cell(prestation_total_row, 1).value = "TOTAL  PRESTATION HT (02)"
        ws.cell(prestation_total_row, 6).number_format = '#,##0.00'
    elif hide_empty_sections():
        ws.delete_rows(prestation_header_row, 3)
    else:
        fill_section(ws, prestation_row, second_section, with_prices=True)


if __name__ == "__main__":
    build(int(sys.argv[1]), sys.argv[2])
