import json
from decimal import Decimal, InvalidOperation
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path


ROOT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_MAPPING_PATH = ROOT_DIR / "resources" / "bpu_st_mapping.json"
NUMBERING_SYSTEMS = {"GENERAL", "ST"}
SOURCE_TYPES = {"GENERAL", "ST", "GENERAL_FALLBACK"}
AUTOMATIC_DESIGNATION_MIN_CONFIDENCE = Decimal("0.82")
BPU_CATEGORY_MARKERS = {"acquisition", "prestation", "fourniture"}


def seed_bpu_st_mapping(connection, mapping_path=None):
    path = Path(mapping_path or DEFAULT_MAPPING_PATH)
    if not path.exists():
        return 0
    payload = json.loads(path.read_text(encoding="utf-8"))
    mappings = payload.get("mappings", [])
    if len(mappings) != payload.get("mapping_count"):
        raise ValueError("Le fichier Mapping BPU ST est incomplet.")
    general_codes = {
        row[0] for row in connection.execute(
            "SELECT article_number FROM bpu_items WHERE is_active=1"
        )
    }
    if not general_codes:
        return 0
    referenced = {int(row["general_article_number"]) for row in mappings}
    expected_catalog_size = payload.get("mapping_count", 0) + payload.get("unmapped_count", 0)
    if len(general_codes) >= expected_catalog_size and not referenced.issubset(general_codes):
        missing = sorted(referenced - general_codes)
        raise ValueError(f"Articles BPU généraux absents du Mapping: {missing[:10]}")
    applicable = [row for row in mappings if int(row["general_article_number"]) in general_codes]
    versioned = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='bpu_st_versions'"
    ).fetchone()
    if versioned:
        version = connection.execute(
            "SELECT id FROM bpu_st_versions WHERE code=?", (payload["version"],)
        ).fetchone()
        if not version:
            cursor = connection.execute(
                """
                INSERT INTO bpu_st_versions(
                    code, status, total_rows, mapped_rows, imported_by,
                    activated_by, activated_at
                ) VALUES(?, 'DRAFT', ?, ?, 'seed', '', NULL)
                """,
                (payload["version"], len(applicable), len(applicable)),
            )
            version_id = cursor.lastrowid
            connection.executemany(
                """
                INSERT INTO bpu_st_items(
                    version_id, st_article_number, general_article_number,
                    confidence, review_status
                ) VALUES(?, ?, ?, 1, 'CONFIRMED')
                """,
                [
                    (version_id, int(row["st_article_number"]), int(row["general_article_number"]))
                    for row in applicable
                ],
            )
            active = connection.execute(
                "SELECT id FROM bpu_st_versions WHERE status='ACTIVE' AND total_rows>0"
            ).fetchone()
            if not active:
                connection.execute(
                    "UPDATE bpu_st_versions SET status='ARCHIVED' WHERE status='ACTIVE'"
                )
                connection.execute(
                    "UPDATE bpu_st_versions SET status='ACTIVE', activated_by='seed', activated_at=CURRENT_TIMESTAMP WHERE id=?",
                    (version_id,),
                )
    else:
        connection.executemany(
            """
            INSERT INTO bpu_st_mapping(st_article_number, general_article_number, mapping_version)
            VALUES(?, ?, ?)
            ON CONFLICT(st_article_number) DO UPDATE SET
                general_article_number=excluded.general_article_number,
                mapping_version=excluded.mapping_version
            """,
            [
                (int(row["st_article_number"]), int(row["general_article_number"]), payload["version"])
                for row in applicable
            ],
        )
        connection.execute(
            f"DELETE FROM bpu_st_mapping WHERE general_article_number NOT IN ({','.join('?' for _ in general_codes)})",
            tuple(general_codes),
        )
    return len(applicable)


def active_mapping_version(connection):
    row = connection.execute(
        "SELECT code FROM bpu_st_versions WHERE status='ACTIVE' LIMIT 1"
    ).fetchone()
    return row[0] if row else ""


def _active_mapping_row(connection, *, st_article_number=None, general_article_number=None):
    versioned = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='bpu_st_items'"
    ).fetchone()
    clauses = ["v.status='ACTIVE'", "i.review_status='CONFIRMED'"]
    params = []
    if st_article_number is not None:
        clauses.append("i.st_article_number=?")
        params.append(st_article_number)
    if general_article_number is not None:
        clauses.append("i.general_article_number=?")
        params.append(general_article_number)
    row = None
    if versioned:
        row = connection.execute(
            f"""
            SELECT i.st_article_number, i.general_article_number, v.code AS mapping_version
            FROM bpu_st_items i
            JOIN bpu_st_versions v ON v.id=i.version_id
            JOIN bpu_items b ON b.article_number=i.general_article_number AND b.is_active=1
            WHERE {' AND '.join(clauses)}
            LIMIT 1
            """,
            tuple(params),
        ).fetchone()
    if row:
        return row
    legacy_clauses = []
    legacy_params = []
    if st_article_number is not None:
        legacy_clauses.append("st_article_number=?")
        legacy_params.append(st_article_number)
    if general_article_number is not None:
        legacy_clauses.append("general_article_number=?")
        legacy_params.append(general_article_number)
    if not legacy_clauses:
        return None
    return connection.execute(
        f"""SELECT m.st_article_number, m.general_article_number, m.mapping_version
            FROM bpu_st_mapping m
            JOIN bpu_items b ON b.article_number=m.general_article_number AND b.is_active=1
            WHERE {' AND '.join('m.' + clause for clause in legacy_clauses)} LIMIT 1""",
        tuple(legacy_params),
    ).fetchone()


def _normalize_text(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def _spreadsheet_rows(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    if path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("Le support XLS nécessite le composant xlrd.") from exc
        workbook = xlrd.open_workbook(path)
        sheet = workbook.sheet_by_index(0)
        return [sheet.row_values(index) for index in range(sheet.nrows)]
    raise ValueError("Format BPU ST non pris en charge. Utilisez XLS ou XLSX.")


def parse_bpu_st_file(path):
    rows = _spreadsheet_rows(path)
    header_index = None
    columns = {}
    for index, row in enumerate(rows[:25]):
        normalized = [_normalize_text(value) for value in row]
        code_index = next((i for i, value in enumerate(normalized) if "code st" in value or value in {"n st", "article st"}), None)
        designation_index = next((i for i, value in enumerate(normalized) if "designation" in value), None)
        if code_index is not None and designation_index is not None:
            header_index = index
            columns = {"code": code_index, "designation": designation_index}
            columns["unite"] = next((i for i, value in enumerate(normalized) if "unite" in value), None)
            columns["price"] = next((i for i, value in enumerate(normalized) if "prix" in value or "pu" in value), None)
            break
    if header_index is None:
        raise ValueError("Colonnes Code ST et Désignation introuvables.")

    parsed = []
    seen = set()
    for row in rows[header_index + 1:]:
        raw_code = row[columns["code"]] if columns["code"] < len(row) else None
        normalized_code = _normalize_text(raw_code)
        if not normalized_code or normalized_code in BPU_CATEGORY_MARKERS:
            continue
        try:
            code = int(float(raw_code))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Code ST invalide: {raw_code}") from exc
        if code in seen:
            raise ValueError(f"Code ST dupliqué: {code}")
        seen.add(code)
        designation = str(row[columns["designation"]] or "").strip() if columns["designation"] < len(row) else ""
        if not designation:
            raise ValueError(f"Désignation manquante pour le code ST {code}.")
        unit_index = columns.get("unite")
        price_index = columns.get("price")
        unit = str(row[unit_index] or "").strip() if unit_index is not None and unit_index < len(row) else ""
        raw_price = row[price_index] if price_index is not None and price_index < len(row) else None
        try:
            price = float(raw_price) if raw_price not in (None, "") else None
        except (TypeError, ValueError):
            price = None
        parsed.append({"st_article_number": code, "designation": designation, "unite": unit, "source_price": price})
    if not parsed:
        raise ValueError("Le fichier BPU ST ne contient aucun article.")
    return parsed


def _designation_similarity(left, right):
    return SequenceMatcher(None, _normalize_text(left), _normalize_text(right)).ratio()


def automatic_designation_mapping(general_rows, st_items):
    """Map ST rows to an ordered one-to-one subset of the general catalogue.

    The score is based on DESIGNATION, with a small unit tie-breaker.  Solving the
    complete ordered sequence avoids locally plausible but globally wrong choices
    when neighboring descriptions are nearly identical.
    """
    general_rows = sorted(
        (dict(row) for row in general_rows),
        key=lambda row: int(row["article_number"]),
    )
    st_items = sorted(
        (dict(row) for row in st_items),
        key=lambda row: int(row["st_article_number"]),
    )
    st_count = len(st_items)
    general_count = len(general_rows)
    if not st_count:
        return {}
    if st_count > general_count:
        raise ValueError(
            "Le BPU ST contient plus d'articles que le BPU général; "
            "le Mapping automatique un-à-un est impossible."
        )

    extra_general_rows = general_count - st_count
    negative = float("-inf")
    scores = [[negative] * (general_count + 1) for _ in range(st_count + 1)]
    choices = [[False] * (general_count + 1) for _ in range(st_count + 1)]
    for column in range(general_count + 1):
        scores[0][column] = 0.0

    for index in range(1, st_count + 1):
        minimum_column = index
        maximum_column = index + extra_general_rows
        st_row = st_items[index - 1]
        st_unit = _normalize_text(st_row.get("unite"))
        for column in range(minimum_column, maximum_column + 1):
            general_row = general_rows[column - 1]
            confidence = _designation_similarity(
                st_row.get("designation"), general_row.get("designation")
            )
            unit_bonus = (
                0.02
                if st_unit and st_unit == _normalize_text(general_row.get("unite"))
                else 0.0
            )
            match_score = scores[index - 1][column - 1] + confidence + unit_bonus
            skip_score = scores[index][column - 1]
            if match_score > skip_score:
                scores[index][column] = match_score
                choices[index][column] = True
            else:
                scores[index][column] = skip_score

    mapped = {}
    index, column = st_count, general_count
    while index:
        if column <= 0:
            raise ValueError("Le Mapping automatique des DESIGNATION a échoué.")
        if choices[index][column]:
            st_row = st_items[index - 1]
            general_row = general_rows[column - 1]
            confidence = _designation_similarity(
                st_row.get("designation"), general_row.get("designation")
            )
            mapped[int(st_row["st_article_number"])] = {
                "general_article_number": int(general_row["article_number"]),
                "confidence": round(confidence, 4),
                "confirmed": Decimal(str(confidence)) >= AUTOMATIC_DESIGNATION_MIN_CONFIDENCE,
            }
            index -= 1
            column -= 1
        else:
            column -= 1
    return mapped


def import_bpu_st_version(connection, path, version_code, actor=""):
    version_code = str(version_code or "").strip()
    if not version_code:
        raise ValueError("Le code de version BPU ST est obligatoire.")
    items = parse_bpu_st_file(path)
    general_rows = [dict(row) for row in connection.execute(
        "SELECT article_number, designation, unite FROM bpu_items WHERE is_active=1 ORDER BY article_number"
    )]
    if not general_rows:
        raise ValueError("Importez d'abord le BPU général.")
    by_number = {int(row["article_number"]): row for row in general_rows}
    automatic_mapping = automatic_designation_mapping(general_rows, items)

    cursor = connection.execute(
        """
        INSERT INTO bpu_st_versions(code, source_filename, status, total_rows, mapped_rows, imported_by)
        VALUES(?, ?, 'DRAFT', ?, 0, ?)
        """,
        (version_code, Path(path).name, len(items), actor),
    )
    version_id = cursor.lastrowid
    mapped = 0
    for item in items:
        automatic = automatic_mapping[item["st_article_number"]]
        candidate = by_number.get(automatic["general_article_number"])
        confidence = float(automatic["confidence"])
        confirmed = bool(candidate and automatic["confirmed"])
        mapped += int(confirmed)
        connection.execute(
            """
            INSERT INTO bpu_st_items(
                version_id, st_article_number, designation, unite, source_price,
                general_article_number, confidence, review_status
            ) VALUES(?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                version_id, item["st_article_number"], item["designation"], item["unite"],
                item["source_price"], int(candidate["article_number"]) if candidate else None,
                round(confidence, 4), "CONFIRMED" if confirmed else "PENDING",
            ),
        )
    connection.execute(
        "UPDATE bpu_st_versions SET mapped_rows=? WHERE id=?", (mapped, version_id)
    )
    return version_id, len(items), mapped


def confirm_bpu_st_mapping(connection, version_id, st_article_number, general_article_number):
    if not connection.execute(
        "SELECT 1 FROM bpu_items WHERE article_number=? AND is_active=1", (general_article_number,)
    ).fetchone():
        raise ValueError("Article BPU ENT introuvable.")
    cursor = connection.execute(
        """
        UPDATE bpu_st_items
        SET general_article_number=?, confidence=1, review_status='CONFIRMED'
        WHERE version_id=? AND st_article_number=?
        """,
        (general_article_number, version_id, st_article_number),
    )
    if not cursor.rowcount:
        raise ValueError("Article BPU ST introuvable.")
    connection.execute(
        """
        UPDATE bpu_st_versions
        SET mapped_rows=(SELECT COUNT(*) FROM bpu_st_items WHERE version_id=? AND review_status='CONFIRMED')
        WHERE id=?
        """,
        (version_id, version_id),
    )


def activate_bpu_st_version(connection, version_id, actor=""):
    version = connection.execute(
        "SELECT * FROM bpu_st_versions WHERE id=?", (version_id,)
    ).fetchone()
    if not version:
        raise ValueError("Version BPU ST introuvable.")
    if int(version["total_rows"]) == 0 or int(version["mapped_rows"]) != int(version["total_rows"]):
        raise ValueError("La couverture du Mapping doit atteindre 100% avant activation.")
    connection.execute("UPDATE bpu_st_versions SET status='ARCHIVED' WHERE status='ACTIVE'")
    connection.execute(
        "UPDATE bpu_st_versions SET status='ACTIVE', activated_by=?, activated_at=CURRENT_TIMESTAMP WHERE id=?",
        (actor, version_id),
    )


def search_bpu_items(connection, query, numbering_system="GENERAL", limit=12):
    query = str(query or "").strip()
    numbering_system = str(numbering_system or "GENERAL").upper()
    if numbering_system not in NUMBERING_SYSTEMS or not query:
        return []
    like = f"%{query}%"
    if numbering_system == "GENERAL":
        rows = connection.execute(
            """
            SELECT b.*, NULL AS st_article_number, 'GENERAL' AS source_reference_type
            FROM bpu_items b
            WHERE b.is_active=1 AND (CAST(b.article_number AS TEXT)=? OR b.designation LIKE ?)
            ORDER BY CASE WHEN CAST(b.article_number AS TEXT)=? THEN 0 ELSE 1 END,
                     b.article_number LIMIT ?
            """,
            (query, like, query, limit),
        ).fetchall()
    else:
        versioned = connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='bpu_st_items'"
        ).fetchone()
        rows = []
        if versioned:
            rows = connection.execute(
                """
                SELECT b.*, i.st_article_number, 'ST' AS source_reference_type
                FROM bpu_st_items i
                JOIN bpu_st_versions v ON v.id=i.version_id AND v.status='ACTIVE'
                JOIN bpu_items b ON b.article_number=i.general_article_number AND b.is_active=1
                WHERE i.review_status='CONFIRMED'
                  AND (CAST(i.st_article_number AS TEXT)=? OR b.designation LIKE ?)
                ORDER BY CASE WHEN CAST(i.st_article_number AS TEXT)=? THEN 0 ELSE 1 END,
                         i.st_article_number LIMIT ?
                """,
                (query, like, query, limit),
            ).fetchall()
        if not rows:
            rows = connection.execute(
                """
                SELECT b.*, m.st_article_number, 'ST' AS source_reference_type
                FROM bpu_st_mapping m
                JOIN bpu_items b ON b.article_number=m.general_article_number AND b.is_active=1
                WHERE CAST(m.st_article_number AS TEXT)=? OR b.designation LIKE ?
                ORDER BY CASE WHEN CAST(m.st_article_number AS TEXT)=? THEN 0 ELSE 1 END,
                         m.st_article_number LIMIT ?
                """,
                (query, like, query, limit),
            ).fetchall()
    return [
        {
            "article_number": row["article_number"],
            "st_article_number": row["st_article_number"],
            "source_reference_type": row["source_reference_type"],
            "designation": row["designation"],
            "unite": row["unite"],
            "pu_ht": row["pu_ht"],
            "categorie": row["categorie"],
        }
        for row in rows
    ]


def resolve_invoice_line(connection, payload, numbering_system="GENERAL"):
    numbering_system = str(numbering_system or "GENERAL").upper()
    if numbering_system not in NUMBERING_SYSTEMS:
        raise ValueError("Système de numérotation invalide.")
    article_number = int(payload["article_number"])
    try:
        quantity = Decimal(str(payload["quantity"]))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise ValueError("La quantité doit être un nombre décimal valide.") from exc
    if quantity <= 0:
        raise ValueError("La quantité doit être positive.")
    default_source = "ST" if numbering_system == "ST" else "GENERAL"
    source_type = str(payload.get("source_reference_type") or default_source).upper()
    source_st = payload.get("source_st_number")
    source_st = int(source_st) if source_st not in (None, "") else None
    if source_type == "GENERAL":
        source_type, source_st = "GENERAL", None
    elif source_type == "ST":
        if source_st is None:
            raise ValueError("La référence BPU ST est obligatoire.")
        mapping = _active_mapping_row(connection, st_article_number=source_st)
        if not mapping or int(mapping["general_article_number"]) != article_number:
            raise ValueError("La référence ST ne correspond pas à l'article BPU général.")
    elif source_type == "GENERAL_FALLBACK":
        if _active_mapping_row(connection, general_article_number=article_number):
            raise ValueError("Cet article possède une référence ST et ne peut pas utiliser le fallback général.")
        source_st = None
    else:
        raise ValueError("Source d'article invalide pour une facture BPU ST.")
    bpu = connection.execute(
        "SELECT * FROM bpu_items WHERE article_number=? AND is_active=1",
        (article_number,),
    ).fetchone()
    if not bpu:
        raise ValueError(f"Article BPU introuvable: {article_number}")
    return bpu, quantity, source_type, source_st


def parse_invoice_line_payload(raw):
    raw = str(raw or "").strip()
    if not raw:
        return []
    if raw.startswith("["):
        payload = json.loads(raw)
        if not isinstance(payload, list):
            raise ValueError("Format des lignes de facture invalide.")
        return payload
    result = []
    for line in raw.replace(";", "\n").splitlines():
        parts = [part.strip() for part in line.replace("\t", ",").split(",") if part.strip()]
        if len(parts) != 2:
            raise ValueError("Chaque ligne article doit être: N article, quantité")
        result.append({"article_number": int(Decimal(parts[0])), "quantity": parts[1]})
    return result
