import json
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.site_partners import site_partner_requirements


PARTNER_HEADERS = ["Raison sociale", "Contact", "Telephone", "E-mail", "Adresse", "Etat"]
PO_HEADERS = [
    "N° BC", "Date BC", "Client", "Direction client", "Direction entreprise",
    "Type BC", "Montant TTC", "Objet",
]
SITE_HEADERS = [
    "N° BC", "Code site", "Nom site", "Sigle typologie", "Libelle complet",
    "Sous-traitant", "BET",
]


def _key(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text.casefold())


def _text(value, maximum=500):
    return " ".join(str(value or "").strip().split())[:maximum]


def _date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    raw = _text(value, 30)
    if not raw:
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError("Date invalide (format attendu: JJ/MM/AAAA).")


def _amount(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).replace("\u00a0", "").replace(" ", "")
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    else:
        raw = raw.replace(",", ".")
    result = float(raw)
    if result < 0:
        raise ValueError("Le montant ne peut pas etre negatif.")
    return result


def _active(value):
    return 0 if _key(value) in {"0", "non", "inactif", "inactive", "desactive"} else 1


def _nature(value):
    aliases = {
        "acq": "ACQUISITION", "acquisition": "ACQUISITION",
        "const": "CONSTRUCTION", "construction": "CONSTRUCTION",
        "constacq": "CONST_ACQUIS", "constaquis": "CONST_ACQUIS",
        "ndc": "NDC", "mgc": "MGC",
    }
    result = aliases.get(_key(value))
    if not result:
        raise ValueError("Type BC invalide.")
    return result


def _sheet_rows(sheet, required_headers):
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"La feuille {sheet.title} est vide.")
    actual = {_key(value): index for index, value in enumerate(values[0]) if _key(value)}
    missing = [header for header in required_headers if _key(header) not in actual]
    if missing:
        raise ValueError("Colonnes manquantes: " + ", ".join(missing))
    result = []
    for source_row, row in enumerate(values[1:], start=2):
        record = {header: row[actual[_key(header)]] if actual[_key(header)] < len(row) else None for header in required_headers}
        if any(value not in (None, "") for value in record.values()):
            record["_row"] = source_row
            result.append(record)
    return result


def _workbook_bytes(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="0B6E4F")
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, start=1):
            width = max(len(str(header)) + 3, *(len(str(row[index - 1] or "")) + 2 for row in rows)) if rows else len(str(header)) + 3
            sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 14), 42)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def partner_template(kind):
    title = "Sous-traitants" if kind == "subcontractors" else "Bureaux_etudes"
    return _workbook_bytes([(title, PARTNER_HEADERS, [])])


def parse_partner_workbook(content, kind, connection):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows = _sheet_rows(workbook.active, PARTNER_HEADERS)
    table = "subcontractors" if kind == "subcontractors" else "design_offices"
    existing = {_key(row["raison_sociale"]): dict(row) for row in connection.execute(f"SELECT * FROM {table}")}
    seen = set()
    preview = []
    for source in rows:
        name = _text(source["Raison sociale"], 160)
        contact = _text(source["Contact"], 160)
        phone = _text(source["Telephone"], 60)
        email = _text(source["E-mail"], 160)
        address = _text(source["Adresse"], 250)
        status, error = "Nouveau", ""
        normalized = _key(name)
        if not name:
            status, error = "Rejete", "Raison sociale obligatoire."
        elif normalized in seen:
            status, error = "Doublon", "Raison sociale repetee dans le fichier."
        elif email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            status, error = "Rejete", "Adresse e-mail invalide."
        elif normalized in existing:
            status = "Mise a jour"
        seen.add(normalized)
        preview.append({
            "row": source["_row"], "raison_sociale": name, "contact": contact,
            "telephone": phone, "email": email, "adresse": address,
            "is_active": _active(source["Etat"]), "status": status, "error": error,
        })
    return {"kind": kind, "rows": preview, "error_count": sum(bool(row["error"]) for row in preview)}


def _lookup(rows, *fields):
    result = {}
    for row in rows:
        for field in fields:
            if field in row.keys() and row[field]:
                result.setdefault(_key(row[field]), row)
    return result


def purchase_order_export(connection):
    po_rows = connection.execute("""
        SELECT po.numero_bc, po.date_bc, c.sigle AS client_sigle,
               cd.sigle AS client_direction_sigle, cb.sigle AS company_direction_sigle,
               po.type_bc, po.montant_ttc, po.objet
        FROM purchase_orders po
        JOIN client_directions cd ON cd.id=po.client_direction_id
        JOIN clients c ON c.id=cd.client_id
        JOIN company_branches cb ON cb.id=po.company_branch_id
        WHERE po.deleted_at IS NULL ORDER BY po.id
    """).fetchall()
    site_rows = connection.execute("""
        SELECT po.numero_bc, s.code_site, s.nom_site, s.typologie_site,
               COALESCE(t.libelle_complet, s.typologie_site) AS typologie_libelle,
               sc.raison_sociale AS subcontractor_name, d.raison_sociale AS design_office_name
        FROM sites s JOIN purchase_orders po ON po.id=s.purchase_order_id
        LEFT JOIN typologies t ON t.id=s.typology_id
        LEFT JOIN subcontractors sc ON sc.id=s.subcontractor_id
        LEFT JOIN design_offices d ON d.id=s.design_office_id
        WHERE s.deleted_at IS NULL AND po.deleted_at IS NULL
        ORDER BY po.id, s.code_site
    """).fetchall()
    pos = [[row["numero_bc"], row["date_bc"], row["client_sigle"], row["client_direction_sigle"],
            row["company_direction_sigle"], row["type_bc"], float(row["montant_ttc"] or 0), row["objet"]] for row in po_rows]
    sites = [[row["numero_bc"], row["code_site"], row["nom_site"], row["typologie_site"],
              row["typologie_libelle"], row["subcontractor_name"] or "", row["design_office_name"] or ""] for row in site_rows]
    return _workbook_bytes([
        ("Bons_de_commande", PO_HEADERS, pos),
        ("Sites", SITE_HEADERS, sites),
    ])


def purchase_order_template():
    return _workbook_bytes([
        ("Bons_de_commande", PO_HEADERS, []),
        ("Sites", SITE_HEADERS, []),
    ])


def parse_purchase_order_workbook(content, connection):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheets = {_key(sheet.title): sheet for sheet in workbook.worksheets}
    po_sheet = sheets.get(_key("Bons_de_commande"))
    site_sheet = sheets.get(_key("Sites"))
    if not po_sheet or not site_sheet:
        raise ValueError("Le fichier doit contenir les feuilles Bons_de_commande et Sites.")
    po_sources = _sheet_rows(po_sheet, PO_HEADERS)
    site_sources = _sheet_rows(site_sheet, SITE_HEADERS)
    clients = _lookup(connection.execute("SELECT * FROM clients WHERE deleted_at IS NULL"), "sigle", "raison_sociale")
    client_directions = list(connection.execute("SELECT * FROM client_directions WHERE deleted_at IS NULL"))
    company_branches = _lookup(connection.execute("SELECT * FROM company_branches"), "sigle", "name")
    subcontractors = _lookup(connection.execute("SELECT * FROM subcontractors"), "raison_sociale")
    design_offices = _lookup(connection.execute("SELECT * FROM design_offices"), "raison_sociale")
    typologies = _lookup(connection.execute("SELECT * FROM typologies"), "sigle")
    existing_pos = {_key(row["numero_bc"]): row for row in connection.execute("SELECT * FROM purchase_orders WHERE deleted_at IS NULL")}
    existing_sites = {_key(row["code_site"]): row for row in connection.execute("SELECT * FROM sites WHERE deleted_at IS NULL")}
    seen_po, parsed_pos = set(), []
    for source in po_sources:
        error = ""
        number = _text(source["N° BC"], 120)
        normalized_number = _key(number)
        try:
            bc_date = _date(source["Date BC"])
            nature = _nature(source["Type BC"])
            amount = _amount(source["Montant TTC"])
        except (ValueError, TypeError) as exc:
            bc_date, nature, amount, error = "", "", 0, str(exc)
        client = clients.get(_key(source["Client"]))
        branch = company_branches.get(_key(source["Direction entreprise"]))
        direction = next((row for row in client_directions if client and row["client_id"] == client["id"] and _key(source["Direction client"]) in {_key(row["sigle"]), _key(row["name"])}), None)
        if not number:
            error = error or "N° BC obligatoire."
        elif normalized_number in seen_po:
            error = error or "N° BC repete dans le fichier."
        elif not client:
            error = error or "Client introuvable."
        elif not direction:
            error = error or "Direction client introuvable pour ce client."
        elif not branch:
            error = error or "Direction entreprise introuvable."
        elif direction["company_branch_id"] != branch["id"]:
            error = error or "La direction client n'est pas rattachee a cette direction entreprise."
        seen_po.add(normalized_number)
        parsed_pos.append({
            "row": source["_row"], "numero_bc": number, "date_bc": bc_date,
            "client_direction_id": direction["id"] if direction else None,
            "company_branch_id": branch["id"] if branch else None,
            "mobilis_direction_id": direction["legacy_mobilis_direction_id"] if direction else None,
            "type_bc": nature, "montant_ttc": amount, "objet": _text(source["Objet"], 500),
            "status": "Mise a jour" if normalized_number in existing_pos else "Nouveau", "error": error,
        })
    po_by_number = {_key(row["numero_bc"]): row for row in parsed_pos}
    seen_sites, parsed_sites = set(), []
    for source in site_sources:
        error = ""
        number = _text(source["N° BC"], 120)
        code = _text(source["Code site"], 120)
        name = _text(source["Nom site"], 200)
        po = po_by_number.get(_key(number))
        typology_sigle = _text(source["Sigle typologie"], 60).upper()
        typology_label = _text(source["Libelle complet"], 500)
        if po and po["type_bc"] in {"MGC", "NDC"}:
            typology_sigle = po["type_bc"]
            known = typologies.get(_key(typology_sigle))
            typology_label = known["libelle_complet"] if known else typology_sigle
        subcontractor = subcontractors.get(_key(source["Sous-traitant"])) if source["Sous-traitant"] else None
        design_office = design_offices.get(_key(source["BET"])) if source["BET"] else None
        if not po:
            error = "BC absent de la feuille Bons_de_commande."
        elif po["error"]:
            error = "BC parent invalide."
        elif not code or not name:
            error = "Code site et nom du site obligatoires."
        elif _key(code) in seen_sites:
            error = "Code site repete dans le fichier."
        elif _key(code) in existing_sites and _key(existing_sites[_key(code)]["purchase_order_id"]) != _key((existing_pos.get(_key(number)) or {"id": ""})["id"]):
            error = "Code site deja rattache a un autre BC."
        elif not typology_sigle or not typology_label:
            error = "Sigle typologie et libelle complet obligatoires."
        else:
            needs_st, needs_bet = site_partner_requirements(po["type_bc"])
            if needs_st and not subcontractor:
                error = "Sous-traitant introuvable ou obligatoire."
            elif needs_bet and not design_office:
                error = "BET introuvable ou obligatoire."
        seen_sites.add(_key(code))
        parsed_sites.append({
            "row": source["_row"], "numero_bc": number, "code_site": code,
            "nom_site": name, "typologie_site": typology_sigle,
            "typologie_libelle": typology_label,
            "subcontractor_id": subcontractor["id"] if subcontractor else None,
            "design_office_id": design_office["id"] if design_office else None,
            "status": "Mise a jour" if _key(code) in existing_sites else "Nouveau", "error": error,
        })
    for po in parsed_pos:
        count = sum(_key(site["numero_bc"]) == _key(po["numero_bc"]) for site in parsed_sites)
        if not po["error"] and count == 0:
            po["error"] = "Aucun site associe."
        elif not po["error"] and po["type_bc"] != "NDC" and count != 1:
            po["error"] = "Un BC hors NDC doit contenir exactement un site."
    errors = sum(bool(row["error"]) for row in parsed_pos + parsed_sites)
    return {"kind": "purchase_orders", "purchase_orders": parsed_pos, "sites": parsed_sites, "error_count": errors}


def save_import_batch(connection, import_type, filename, payload, actor):
    token = uuid4().hex
    connection.execute(
        "INSERT INTO data_import_batches(token,import_type,source_filename,payload_json,error_count,created_by) VALUES(?,?,?,?,?,?)",
        (token, import_type, filename, json.dumps(payload, ensure_ascii=False), payload["error_count"], actor),
    )
    return token


def load_import_batch(connection, token, import_type):
    row = connection.execute(
        """SELECT * FROM data_import_batches
           WHERE token=? AND import_type=? AND status='preview'
             AND created_at >= datetime('now', '-24 hours')""",
        (token, import_type),
    ).fetchone()
    if not row:
        raise ValueError("Apercu d'import introuvable ou deja utilise.")
    return row, json.loads(row["payload_json"])


def apply_partner_import(connection, batch, payload, actor):
    if payload["error_count"]:
        raise ValueError("Le fichier contient des erreurs. Corrigez-le avant de confirmer.")
    table = "subcontractors" if payload["kind"] == "subcontractors" else "design_offices"
    existing_by_name = {
        _key(row["raison_sociale"]): row["id"]
        for row in connection.execute(f"SELECT id,raison_sociale FROM {table}")
    }
    for row in payload["rows"]:
        existing_id = existing_by_name.get(_key(row["raison_sociale"]))
        if existing_id:
            connection.execute(
                f"UPDATE {table} SET raison_sociale=?,contact=?,telephone=?,email=?,adresse=?,is_active=?,updated_by=? WHERE id=?",
                (row["raison_sociale"], row["contact"], row["telephone"], row["email"], row["adresse"], row["is_active"], actor, existing_id),
            )
        else:
            hidden_sigle = f"REF-{uuid4().hex[:12].upper()}"
            connection.execute(
                f"INSERT INTO {table}(raison_sociale,sigle,contact,telephone,email,adresse,is_active,created_by,updated_by) VALUES(?,?,?,?,?,?,?,?,?)",
                (row["raison_sociale"], hidden_sigle, row["contact"], row["telephone"], row["email"], row["adresse"], row["is_active"], actor, actor),
            )
            existing_by_name[_key(row["raison_sociale"])] = connection.execute("SELECT last_insert_rowid()").fetchone()[0]
    connection.execute(
        "UPDATE data_import_batches SET status='applied',applied_by=?,applied_at=CURRENT_TIMESTAMP WHERE id=?",
        (actor, batch["id"]),
    )


def apply_purchase_order_import(connection, batch, payload, actor):
    if payload["error_count"]:
        raise ValueError("Le fichier contient des erreurs. Corrigez-le avant de confirmer.")
    po_ids = {}
    for row in payload["purchase_orders"]:
        existing = connection.execute("SELECT id FROM purchase_orders WHERE numero_bc=? AND deleted_at IS NULL", (row["numero_bc"],)).fetchone()
        if existing:
            po_id = existing["id"]
            connection.execute("""
                UPDATE purchase_orders SET date_bc=?,mobilis_direction_id=?,client_direction_id=?,company_branch_id=?,
                    type_bc=?,objet=?,montant_ttc=?,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?
            """, (row["date_bc"] or None, row["mobilis_direction_id"], row["client_direction_id"], row["company_branch_id"],
                  row["type_bc"], row["objet"], row["montant_ttc"], actor, po_id))
        else:
            po_id = connection.execute("""
                INSERT INTO purchase_orders(numero_bc,date_bc,mobilis_direction_id,client_direction_id,company_branch_id,
                    type_bc,objet,montant_ttc,created_by,updated_by) VALUES(?,?,?,?,?,?,?,?,?,?)
            """, (row["numero_bc"], row["date_bc"] or None, row["mobilis_direction_id"], row["client_direction_id"],
                  row["company_branch_id"], row["type_bc"], row["objet"], row["montant_ttc"], actor, actor)).lastrowid
        po_ids[_key(row["numero_bc"])] = po_id
    for row in payload["sites"]:
        po_id = po_ids[_key(row["numero_bc"])]
        typology = connection.execute("SELECT id FROM typologies WHERE sigle=? COLLATE NOCASE", (row["typologie_site"],)).fetchone()
        if typology:
            typology_id = typology["id"]
            connection.execute("UPDATE typologies SET libelle_complet=?,is_active=1,updated_by=? WHERE id=?", (row["typologie_libelle"], actor, typology_id))
        else:
            typology_id = connection.execute(
                "INSERT INTO typologies(sigle,libelle_complet,created_by,updated_by) VALUES(?,?,?,?)",
                (row["typologie_site"], row["typologie_libelle"], actor, actor),
            ).lastrowid
        existing = connection.execute("SELECT id FROM sites WHERE code_site=? AND deleted_at IS NULL", (row["code_site"],)).fetchone()
        if existing:
            connection.execute("""
                UPDATE sites SET purchase_order_id=?,nom_site=?,region='',typologie_site=?,typology_id=?,bet='',
                    subcontractor_id=?,design_office_id=?,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?
            """, (po_id, row["nom_site"], row["typologie_site"], typology_id, row["subcontractor_id"], row["design_office_id"], actor, existing["id"]))
        else:
            connection.execute("""
                INSERT INTO sites(purchase_order_id,code_site,nom_site,region,typologie_site,typology_id,bet,
                    subcontractor_id,design_office_id,created_by,updated_by) VALUES(?,?,?,'',?,?,'',?,?,?,?)
            """, (po_id, row["code_site"], row["nom_site"], row["typologie_site"], typology_id,
                  row["subcontractor_id"], row["design_office_id"], actor, actor))
    connection.execute(
        "UPDATE data_import_batches SET status='applied',applied_by=?,applied_at=CURRENT_TIMESTAMP WHERE id=?",
        (actor, batch["id"]),
    )


def import_report(payload):
    if payload["kind"] == "purchase_orders":
        po_rows = [[row["row"], row["numero_bc"], row["status"], row["error"]] for row in payload["purchase_orders"]]
        site_rows = [[row["row"], row["numero_bc"], row["code_site"], row["status"], row["error"]] for row in payload["sites"]]
        return _workbook_bytes([
            ("Bons_de_commande", ["Ligne", "N° BC", "Statut", "Erreur"], po_rows),
            ("Sites", ["Ligne", "N° BC", "Code site", "Statut", "Erreur"], site_rows),
        ])
    rows = [[row["row"], row["raison_sociale"], row["status"], row["error"]] for row in payload["rows"]]
    return _workbook_bytes([("Rapport", ["Ligne", "Raison sociale", "Statut", "Erreur"], rows)])
