import sqlite3
from io import BytesIO

import pytest
from openpyxl import load_workbook

from database.migrations import run_migrations
from test_migrations import legacy_database
from services.data_imports import (
    apply_partner_import,
    apply_purchase_order_import,
    import_report,
    parse_partner_workbook,
    parse_purchase_order_workbook,
    partner_template,
    purchase_order_export,
    purchase_order_template,
    save_import_batch,
)


@pytest.fixture
def migrated_connection(tmp_path):
    database_path = tmp_path / "imports.sqlite3"
    connection, *_ = legacy_database(database_path)
    try:
        run_migrations(connection)
        yield connection
    finally:
        connection.close()


def workbook_with_rows(content, sheet_rows):
    workbook = load_workbook(BytesIO(content))
    for sheet_name, rows in sheet_rows.items():
        sheet = workbook[sheet_name]
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.integration
def test_partner_template_preview_apply_and_report(migrated_connection):
    content = workbook_with_rows(
        partner_template("subcontractors"),
        {"Sous-traitants": [["SARL Atlas", "Ali Benali", "0550000000", "ali@atlas.dz", "Oran", "Actif"]]},
    )
    payload = parse_partner_workbook(content, "subcontractors", migrated_connection)
    assert payload["error_count"] == 0
    assert payload["rows"][0]["status"] == "Nouveau"
    token = save_import_batch(migrated_connection, "subcontractors", "st.xlsx", payload, "admin")
    batch = migrated_connection.execute("SELECT * FROM data_import_batches WHERE token=?", (token,)).fetchone()
    apply_partner_import(migrated_connection, batch, payload, "admin")
    row = migrated_connection.execute(
        "SELECT raison_sociale,contact,email FROM subcontractors WHERE raison_sociale='SARL Atlas'"
    ).fetchone()
    assert tuple(row) == ("SARL Atlas", "Ali Benali", "ali@atlas.dz")
    assert load_workbook(BytesIO(import_report(payload))).active["B2"].value == "SARL Atlas"


@pytest.mark.integration
def test_partner_import_rejects_duplicates(migrated_connection):
    content = workbook_with_rows(
        partner_template("design_offices"),
        {"Bureaux_etudes": [["BET Centre", "Contact 1", "", "", "", "Actif"], [" bet  centre ", "Contact 2", "", "", "", "Actif"]]},
    )
    payload = parse_partner_workbook(content, "design_offices", migrated_connection)
    assert payload["error_count"] == 1
    assert payload["rows"][1]["status"] == "Doublon"


@pytest.mark.integration
def test_purchase_order_import_and_export_round_trip(migrated_connection):
    client = migrated_connection.execute("SELECT * FROM clients WHERE deleted_at IS NULL LIMIT 1").fetchone()
    direction = migrated_connection.execute("SELECT * FROM client_directions WHERE deleted_at IS NULL LIMIT 1").fetchone()
    branch = migrated_connection.execute("SELECT * FROM company_branches LIMIT 1").fetchone()
    migrated_connection.execute(
        "INSERT INTO subcontractors(raison_sociale,sigle,contact) VALUES('SARL Atlas','REF-ST','Ali')"
    )
    content = workbook_with_rows(
        purchase_order_template(),
        {
            "Bons_de_commande": [["BC-IMPORT-001", "24/08/2026", client["sigle"], direction["sigle"], branch["sigle"], "CONST", 120000, "Travaux"]],
            "Sites": [["BC-IMPORT-001", "SITE-IMPORT-001", "Site importé", "A12", "A12 ( MAT 12M + BTS OUTDOOR )", "SARL Atlas", ""]],
        },
    )
    payload = parse_purchase_order_workbook(content, migrated_connection)
    assert payload["error_count"] == 0
    token = save_import_batch(migrated_connection, "purchase_orders", "bc.xlsx", payload, "admin")
    batch = migrated_connection.execute("SELECT * FROM data_import_batches WHERE token=?", (token,)).fetchone()
    apply_purchase_order_import(migrated_connection, batch, payload, "admin")
    assert migrated_connection.execute("SELECT COUNT(*) FROM purchase_orders WHERE numero_bc='BC-IMPORT-001'").fetchone()[0] == 1
    site = migrated_connection.execute(
        "SELECT sc.raison_sociale FROM sites s JOIN subcontractors sc ON sc.id=s.subcontractor_id WHERE s.code_site='SITE-IMPORT-001'"
    ).fetchone()
    assert site[0] == "SARL Atlas"
    exported = load_workbook(BytesIO(purchase_order_export(migrated_connection)), data_only=True)
    assert {sheet.title for sheet in exported.worksheets} == {"Bons_de_commande", "Sites"}
    assert any(row[0] == "BC-IMPORT-001" for row in exported["Bons_de_commande"].iter_rows(min_row=2, values_only=True))


@pytest.mark.integration
def test_purchase_order_import_rejects_unknown_references(migrated_connection):
    content = workbook_with_rows(
        purchase_order_template(),
        {
            "Bons_de_commande": [["BC-INVALID", "24/08/2026", "CLIENT-X", "DR-X", "DRO", "NDC", 0, ""]],
            "Sites": [["BC-INVALID", "SITE-X", "Site X", "NDC", "NDC", "", "BET inconnu"]],
        },
    )
    payload = parse_purchase_order_workbook(content, migrated_connection)
    assert payload["error_count"] >= 2
