import json
import re
import threading
from io import BytesIO
from http.cookiejar import CookieJar
from urllib.parse import parse_qs, urlencode, urlparse
from urllib.error import HTTPError
from urllib.request import HTTPCookieProcessor, Request, build_opener

import db as db_module
import app as app_module
import pytest
from openpyxl import load_workbook
import services.backups as backups
import services.licensing as licensing
from app import APP_VERSION, App
from http.server import ThreadingHTTPServer
from services.xlsx import make_xlsx
from services.machine_identity import verify_activation_request
from services.data_imports import partner_template, purchase_order_template


pytestmark = pytest.mark.e2e


@pytest.fixture(autouse=True)
def isolated_license_file(tmp_path, monkeypatch):
    monkeypatch.setattr(licensing, "LICENSE_PATH", tmp_path / "license.json")


def _csrf(html):
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    if not match:
        match = re.search(r'name="csrf-token" content="([^"]+)"', html)
    assert match, html[:500]
    return match.group(1)


def _request(opener, url, data=None, headers=None, timeout=30):
    payload = None
    if isinstance(data, dict):
        payload = urlencode(data).encode("utf-8")
    elif data is not None:
        payload = data
    request = Request(url, data=payload, headers=headers or {})
    try:
        return opener.open(request, timeout=timeout)
    except HTTPError as exc:
        return exc


def _read(response):
    try:
        return response.read()
    finally:
        response.close()


def _multipart(fields, files):
    boundary = "----phoenix-e2e-boundary"
    chunks = []
    for name, value in fields.items():
        chunks.extend([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
            str(value).encode(),
            b"\r\n",
        ])
    for name, file_info in files.items():
        filename, content = file_info
        chunks.extend([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="{name}"; filename="{filename}"\r\n'.encode(),
            b"Content-Type: application/octet-stream\r\n\r\n",
            content,
            b"\r\n",
        ])
    chunks.append(f"--{boundary}--\r\n".encode())
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


def test_first_run_csrf_backup_restore_e2e(tmp_path, monkeypatch):
    database_path = tmp_path / "data" / "e2e.sqlite3"
    backup_dir = tmp_path / "backups"
    monkeypatch.setenv("PHOENIX_DB_PATH", str(database_path))
    monkeypatch.setattr(db_module, "DB_PATH", database_path)
    monkeypatch.setattr(backups, "DB_PATH", database_path)
    monkeypatch.setattr(backups, "BACKUP_DIR", backup_dir)

    server = ThreadingHTTPServer(("127.0.0.1", 0), App)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    opener = build_opener(HTTPCookieProcessor(CookieJar()))

    try:
        setup_html = _read(_request(opener, base_url + "/setup")).decode("utf-8")
        setup_token = _csrf(setup_html)
        response = _request(opener, base_url + "/setup", {
            "csrf_token": setup_token,
            "new_password": "Market123",
            "confirm_password": "Market123",
        })
        assert response.status == 200
        assert "Tableau" in _read(response).decode("utf-8")

        login_with_existing_session = _read(
            _request(opener, base_url + "/login")
        ).decode("utf-8")
        assert login_with_existing_session.count('name="csrf_token"') == 1
        assert 'autocomplete="username"' in login_with_existing_session
        assert 'autocomplete="current-password"' in login_with_existing_session
        assert '<body class="auth-page">' in login_with_existing_session
        assert 'class="auth-card"' in login_with_existing_session
        assert '/static/app-ui.js' in login_with_existing_session
        relogin = _request(opener, base_url + "/login", {
            "csrf_token": _csrf(login_with_existing_session),
            "username": "admin",
            "password": "Market123",
        })
        assert relogin.status == 200
        for header, expected in {
            "X-Content-Type-Options": "nosniff",
            "X-Frame-Options": "DENY",
            "Referrer-Policy": "same-origin",
            "Cache-Control": "no-store",
        }.items():
            assert relogin.headers[header] == expected
        assert "frame-ancestors 'none'" in relogin.headers["Content-Security-Policy"]
        assert "Tableau" in _read(relogin).decode("utf-8")

        for state_path in (
            "/logout",
            "/mobilis/delete?id=1",
            "/purchase-orders/delete?id=1",
            "/purchase-orders/document/delete?id=1",
            "/purchase-orders/site/delete?id=1&po_id=1",
            "/sites/delete?id=1",
            "/invoices/delete?id=1",
            "/archives/delete?file=test.pdf",
        ):
            rejected_get = _request(opener, base_url + state_path)
            assert rejected_get.status == 405
            assert rejected_get.headers["Allow"] == "POST"
            _read(rejected_get)

        users_page = _request(opener, base_url + "/users")
        assert users_page.status == 200
        users_html = _read(users_page).decode("utf-8")
        assert "Utilisateurs" in users_html
        user_token = _csrf(users_html)
        create_user = _request(opener, base_url + "/users/create", {
            "csrf_token": user_token,
            "username": "viewer_e2e",
            "password": "Viewer123",
            "role": "viewer",
            "company_branch_id": "1",
        })
        assert create_user.status == 200
        assert "viewer_e2e" in _read(create_user).decode("utf-8")

        settings_page = _request(opener, base_url + "/settings")
        assert settings_page.status == 200
        settings_html = _read(settings_page).decode("utf-8")
        assert "Parametres" in settings_html
        assert "/users" in settings_html
        assert 'data-path="/users"' not in settings_html
        assert 'data-path="/settings"' in settings_html

        about_page = _request(opener, base_url + "/about")
        assert about_page.status == 200
        assert APP_VERSION in _read(about_page).decode("utf-8")

        license_page = _request(opener, base_url + "/license")
        assert license_page.status == 200
        license_html = _read(license_page).decode("utf-8")
        assert "demo" in license_html.lower()
        assert "/license/request" in license_html
        activation_download = _request(opener, base_url + "/license/request")
        assert activation_download.status == 200
        assert "attachment" in activation_download.headers.get("Content-Disposition", "")
        activation_document = json.loads(_read(activation_download).decode("utf-8"))
        assert verify_activation_request(activation_document)["machine_id"].startswith("PHX-")

        guide_pdf = _request(opener, base_url + "/docs/Guide_utilisateur_PhoEniX_BPU.pdf")
        assert guide_pdf.status == 200
        assert _read(guide_pdf).startswith(b"%PDF")

        with db_module.db() as con:
            con.execute(
                "UPDATE company_settings SET nom='SAPTA', rgc='RGC', nif='NIF', art='ART', adresse='Adresse', numero_compte='RIB' WHERE id=1"
            )
            con.execute("UPDATE clients SET reference_contrat='CTR-001' WHERE id=1")
            direction_id = con.execute(
                "INSERT INTO mobilis_directions(doit_nom, direction_regionale, adresse, rgc, nif) VALUES('Mobilis', 'DR Alger', 'Alger', 'RGC-M', 'NIF-M')"
            ).lastrowid
            con.execute(
                "INSERT INTO bpu_items(article_number, designation, unite, pu_ht, categorie) VALUES(1, 'Cable fibre', 'm', 100, 'fourniture')"
            )
            po_id = con.execute(
                "INSERT INTO purchase_orders(numero_bc, date_bc, mobilis_direction_id, type_bc, objet, montant_ttc) VALUES('BC-E2E', '2026-08-09', ?, 'CONSTRUCTION', 'Projet E2E', 1190)",
                (direction_id,),
            ).lastrowid
            site_id = con.execute(
                "INSERT INTO sites(purchase_order_id, code_site, nom_site, region, typologie_site) VALUES(?, 'SITE-E2E', 'Site E2E', 'Alger', 'A12')",
                (po_id,),
            ).lastrowid
            invoice_id = con.execute(
                """
                INSERT INTO invoices(
                    invoice_number, invoice_type, purchase_order_id, site_id, invoice_date,
                    total_ht, retenue_garantie, montant_ht_apres_rg, tva, total_ttc, montant_en_lettres
                ) VALUES('FAC-E2E', 'CONSTRUCTION', ?, ?, '2026-08-09', 1000, 50, 950, 180.5, 1130.5, 'MILLE CENT TRENTE')
                """,
                (po_id, site_id),
            ).lastrowid
            con.execute(
                """
                INSERT INTO invoice_lines(
                    invoice_id, article_number, designation_snapshot, unite_snapshot,
                    pu_ht_snapshot, categorie_snapshot, quantite, montant_ht
                ) VALUES(?, 1, 'Cable fibre', 'm', 100, 'fourniture', 10, 1000)
                """,
                (invoice_id,),
            )

        viewer_opener = build_opener(HTTPCookieProcessor(CookieJar()))
        viewer_login_html = _read(_request(viewer_opener, base_url + "/login")).decode("utf-8")
        viewer_login = _request(viewer_opener, base_url + "/login", {
            "csrf_token": _csrf(viewer_login_html),
            "username": "viewer_e2e",
            "password": "Viewer123",
        })
        assert viewer_login.status == 200
        _read(viewer_login)
        viewer_invoices = _request(viewer_opener, base_url + "/invoices")
        assert viewer_invoices.status == 200
        viewer_invoices_html = _read(viewer_invoices).decode("utf-8")
        viewer_export = _request(
            viewer_opener,
            base_url + f"/invoices/export?id={invoice_id}",
        )
        assert viewer_export.status == 403
        _read(viewer_export)
        viewer_create = _request(viewer_opener, base_url + "/invoices", {
            "csrf_token": _csrf(viewer_invoices_html),
            "invoice_number": "FORBIDDEN",
        })
        assert viewer_create.status == 403
        _read(viewer_create)
        viewer_delete = _request(viewer_opener, base_url + "/invoices/delete", {
            "csrf_token": _csrf(viewer_invoices_html),
            "id": str(invoice_id),
        })
        assert viewer_delete.status == 403
        _read(viewer_delete)

        invoices_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
        assert 'id="invoice-site-select"' in invoices_html
        assert 'data-placeholder="Saisir ou rechercher un code site"' in invoices_html
        assert "Factures enregistrées" not in invoices_html
        assert "Afficher les factures annulées" not in invoices_html
        invoice_token = _csrf(invoices_html)
        empty_without_po_response = _request(opener, base_url + "/invoices", {
            "csrf_token": invoice_token,
            "invoice_number": "EMPTY-WITHOUT-PO-E2E",
            "invoice_date": "",
            "purchase_order_id": "",
            "site_id": "",
            "lines": "",
        })
        assert empty_without_po_response.status == 200
        assert "Sélectionnez un bon de commande ou ajoutez au moins un article" in _read(
            empty_without_po_response
        ).decode("utf-8")

        invoices_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
        articles_without_po_response = _request(opener, base_url + "/invoices", {
            "csrf_token": _csrf(invoices_html),
            "invoice_number": "DRAFT-E2E",
            "invoice_date": "",
            "purchase_order_id": "",
            "site_id": str(site_id),
            "lines": "1,2",
        })
        assert articles_without_po_response.status == 200
        assert "Brouillon de facture" in _read(articles_without_po_response).decode("utf-8")

        with db_module.db() as con:
            assert con.execute(
                "SELECT COUNT(*) FROM invoices WHERE invoice_number='EMPTY-WITHOUT-PO-E2E'"
            ).fetchone()[0] == 0
            article_draft = con.execute(
                """
                SELECT id, purchase_order_id, site_id, invoice_type, total_ht, lifecycle_status
                FROM invoice_lifecycle WHERE invoice_number='DRAFT-E2E'
                """
            ).fetchone()
            assert tuple(article_draft)[1:] == (None, site_id, None, 200, "BROUILLON")
            assert tuple(con.execute(
                """
                SELECT article_number, quantite, montant_ht
                FROM invoice_lines WHERE invoice_id=?
                """,
                (article_draft["id"],),
            ).fetchone()) == (1, 2, 200)

        draft_tracking_html = _read(
            _request(opener, base_url + "/table-facturation-new")
        ).decode("utf-8")
        assert "DRAFT-E2E" in draft_tracking_html
        assert "SITE-E2E" in draft_tracking_html
        assert "Brouillon de facture" in draft_tracking_html

        with db_module.db() as con:
            subcontractor_id = con.execute(
                "INSERT INTO subcontractors(raison_sociale,sigle) VALUES('Sous-traitant E2E','ST E2E')"
            ).lastrowid
            mgc_po_id = con.execute(
                """
                INSERT INTO purchase_orders(
                    numero_bc, date_bc, mobilis_direction_id, type_bc, objet, montant_ttc
                ) VALUES('BC-MGC-E2E', '2026-08-10', ?, 'MGC', 'Maintenance', 2261)
                """,
                (direction_id,),
            ).lastrowid
            mgc_site_id = con.execute(
                """
                INSERT INTO sites(purchase_order_id, code_site, nom_site, subcontractor_id)
                VALUES(?, 'SITE-MGC-E2E', 'Site MGC E2E', ?)
                """,
                (mgc_po_id, subcontractor_id),
            ).lastrowid
            replacement_po_id = con.execute(
                """
                INSERT INTO purchase_orders(
                    numero_bc, date_bc, mobilis_direction_id, type_bc, objet, montant_ttc
                ) VALUES('BC-REPLACE-E2E', '2026-08-10', ?, 'CONSTRUCTION',
                         'Remplacement du site', 226)
                """,
                (direction_id,),
            ).lastrowid
            replacement_site_id = con.execute(
                """
                INSERT INTO sites(
                    purchase_order_id, code_site, nom_site, region, typologie_site,
                    subcontractor_id
                ) VALUES(?, 'SITE-REPLACE-E2E', 'Site remplacement E2E', 'Oran', 'A12', ?)
                """,
                (replacement_po_id, subcontractor_id),
            ).lastrowid

        invoices_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
        incomplete_with_po_response = _request(opener, base_url + "/invoices", {
            "csrf_token": _csrf(invoices_html),
            "invoice_number": "DRAFT-BC-E2E",
            "invoice_date": "",
            "purchase_order_id": str(mgc_po_id),
            "site_id": str(mgc_site_id),
            "lines": "",
        })
        assert incomplete_with_po_response.status == 200
        assert "Brouillon de facture" in _read(incomplete_with_po_response).decode("utf-8")
        with db_module.db() as con:
            mgc_draft = con.execute(
                """
                SELECT id, lifecycle_status FROM invoice_lifecycle
                WHERE invoice_number='DRAFT-BC-E2E'
                """
            ).fetchone()
            assert mgc_draft["lifecycle_status"] == "BROUILLON"
            mgc_draft_id = mgc_draft["id"]

        invoices_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
        mgc_response = _request(opener, base_url + "/invoices", {
            "csrf_token": _csrf(invoices_html),
            "id": str(mgc_draft_id),
            "invoice_number": "FAC-MGC-E2E",
            "invoice_date": "2026-08-10",
            "purchase_order_id": str(mgc_po_id),
            "site_id": str(mgc_site_id),
            "lines": "1,10",
        })
        assert mgc_response.status == 200
        assert "Prête pour dépôt DTC" in _read(mgc_response).decode("utf-8")

        invoices_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
        assert "Le site précédent ne correspond pas au BC sélectionné" in invoices_html
        replace_site_response = _request(opener, base_url + "/invoices", {
            "csrf_token": _csrf(invoices_html),
            "id": str(article_draft["id"]),
            "invoice_number": "DRAFT-E2E",
            "invoice_date": "",
            "purchase_order_id": str(replacement_po_id),
            "site_id": str(site_id),
            "lines": "1,2",
        })
        assert replace_site_response.status == 200
        replace_site_html = _read(replace_site_response).decode("utf-8")
        replace_site_alert = re.search(r'<section class="alert">([^<]+)</section>', replace_site_html)
        assert "Brouillon de facture" in replace_site_html, (
            replace_site_alert.group(1) if replace_site_alert else "Aucun message"
        )

        with db_module.db() as con:
            mgc_invoice = con.execute(
                """
                SELECT id, invoice_type, lifecycle_status
                FROM invoice_lifecycle WHERE invoice_number='FAC-MGC-E2E'
                """
            ).fetchone()
            assert mgc_invoice["invoice_type"] == "CONSTRUCTION"
            assert mgc_invoice["lifecycle_status"] == "READY_DTC"
            draft = con.execute(
                """
                SELECT purchase_order_id, site_id, lifecycle_status
                FROM invoice_lifecycle WHERE invoice_number='DRAFT-E2E'
                """
            ).fetchone()
            assert draft["purchase_order_id"] == replacement_po_id
            assert draft["site_id"] == replacement_site_id
            assert draft["lifecycle_status"] == "BROUILLON"
            legacy_ndc_id = con.execute(
                """
                INSERT INTO invoices(
                    invoice_number, invoice_type, purchase_order_id, created_by, updated_by
                ) VALUES('NDC-LEGACY-E2E', 'NDC', ?, 'migration', 'migration')
                """,
                (po_id,),
            ).lastrowid
            con.executemany(
                "INSERT INTO invoice_sites(invoice_id, site_id) VALUES(?, ?)",
                ((legacy_ndc_id, site_id), (legacy_ndc_id, mgc_site_id)),
            )

        legacy_ndc_html = _read(
            _request(opener, base_url + f"/invoices?edit_id={legacy_ndc_id}")
        ).decode("utf-8")
        assert "Hors BC historique" in legacy_ndc_html
        assert legacy_ndc_html.count('class="check-row"') >= 2

        tracking_html = _read(
            _request(opener, base_url + "/table-facturation-new")
        ).decode("utf-8")
        assert "DR MOBILIS" in tracking_html
        assert "OBJET BC" in tracking_html
        assert "DATE DÉPÔT DTC" in tracking_html
        assert "DATE DÉPÔT MOBILIS" in tracking_html
        assert "DATE PAIEMENT" in tracking_html
        assert "N° ORDRE DE VIREMENT" in tracking_html
        assert "Afficher les factures annulées" in tracking_html
        assert 'class="facturation-title-row"' in tracking_html
        assert 'name="company_branch" form="facturation-filter-form"' in tracking_html
        assert 'name="client" form="facturation-filter-form"' in tracking_html
        assert 'title="Exporter Excel" aria-label="Exporter Excel"' in tracking_html
        assert ">Exporter Excel</span>" not in tracking_html
        assert 'class="ndc-sites-details"' in tracking_html
        ndc_popup = re.search(
            r'<details class="ndc-sites-details"><summary>2 sites</summary><ul>(.*?)</ul>',
            tracking_html,
        )
        assert ndc_popup
        assert re.findall(r"<li>(.*?)</li>", ndc_popup.group(1)) == [
            "SITE-E2E",
            "SITE-MGC-E2E",
        ]
        assert "A12," not in tracking_html
        assert 'name="date_start" value=""' in tracking_html
        assert 'name="date_end" value=""' in tracking_html
        assert "DÉPOSÉ</span>" not in tracking_html
        tracking_token = _csrf(invoices_html)
        for field, value, expected_status in (
            ("date_depot_dtc", "2026-08-11", "DEPOSITED_DTC"),
            ("date_depot_mobilis", "2026-08-12", "AWAITING_PAYMENT"),
            ("date_ov", "2026-08-13", "PAID"),
        ):
            tracking_response = _request(
                opener,
                base_url + "/table-facturation-new/update",
                {
                    "csrf_token": tracking_token,
                    "invoice_id": str(mgc_invoice["id"]),
                    field: value,
                },
            )
            assert tracking_response.status == 200
            assert expected_status in _read(tracking_response).decode("utf-8")

        payment_reference_response = _request(
            opener,
            base_url + "/table-facturation-new/update",
            {
                "csrf_token": tracking_token,
                "invoice_id": str(mgc_invoice["id"]),
                "numero_ordre_virement": "OV-E2E-2026-001",
            },
        )
        assert payment_reference_response.status == 200
        payment_reference_payload = json.loads(
            _read(payment_reference_response).decode("utf-8")
        )
        assert payment_reference_payload["numero_ordre_virement"] == "OV-E2E-2026-001"

        locked_edit = _request(opener, base_url + "/invoices", {
            "csrf_token": tracking_token,
            "id": str(mgc_invoice["id"]),
            "invoice_number": "FAC-MGC-MODIFIED",
            "invoice_date": "2026-08-10",
            "purchase_order_id": str(mgc_po_id),
            "site_id": str(mgc_site_id),
            "lines": "1,1",
        })
        assert locked_edit.status == 200
        assert "verrouillées" in _read(locked_edit).decode("utf-8")
        with db_module.db() as con:
            locked_number = con.execute(
                "SELECT invoice_number FROM invoices WHERE id=?", (mgc_invoice["id"],)
            ).fetchone()[0]
            assert locked_number == "FAC-MGC-E2E"

        cancelled = _request(opener, base_url + "/table-facturation-new/update", {
            "csrf_token": tracking_token,
            "invoice_id": str(mgc_invoice["id"]),
            "action": "cancel",
            "reason": "Test affichage des factures annulées",
        })
        assert cancelled.status == 200
        cancelled_payload = json.loads(_read(cancelled).decode("utf-8"))
        assert cancelled_payload["ok"] is True
        assert cancelled_payload["status"] == "CANCELLED"

        default_tracking = _read(
            _request(opener, base_url + "/table-facturation-new")
        ).decode("utf-8")
        assert "FAC-MGC-E2E" not in default_tracking
        cancelled_tracking = _read(
            _request(opener, base_url + "/table-facturation-new?status=CANCELLED")
        ).decode("utf-8")
        assert "FAC-MGC-E2E" in cancelled_tracking
        assert 'data-lifecycle-action="restore"' in cancelled_tracking
        assert 'title="Restaurer la facture"' in cancelled_tracking

        cancelled_toggle_tracking = _read(
            _request(opener, base_url + "/table-facturation-new?show_cancelled=1")
        ).decode("utf-8")
        assert "FAC-MGC-E2E" in cancelled_toggle_tracking
        assert "DRAFT-E2E" in cancelled_toggle_tracking
        assert "NDC-LEGACY-E2E" in cancelled_toggle_tracking
        assert "Annulée" in cancelled_toggle_tracking
        assert 'name="show_cancelled" value="1" checked' in cancelled_toggle_tracking

        invoice_creation_page = _read(
            _request(opener, base_url + "/invoices")
        ).decode("utf-8")
        assert "Factures enregistrées" not in invoice_creation_page
        assert "Afficher les factures annulées" not in invoice_creation_page
        cancelled_edit = _read(
            _request(
                opener,
                base_url + f'/invoices?edit_id={mgc_invoice["id"]}',
            )
        ).decode("utf-8")
        assert "Facture annulée : les données financières sont verrouillées" in cancelled_edit
        assert '<fieldset class="invoice-form-fields" disabled>' in cancelled_edit

        excel = _request(opener, base_url + f"/invoices/export?id={invoice_id}")
        assert excel.status == 200
        assert excel.headers["Content-Type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert excel.headers["X-Content-Type-Options"] == "nosniff"
        assert _read(excel).startswith(b"PK")
        with db_module.db() as con:
            assert con.execute(
                """
                SELECT COUNT(*) FROM invoice_tracking_events
                WHERE invoice_id=? AND event_type='INVOICE_EXPORTED'
                  AND new_value='invoice_xlsx'
                """,
                (invoice_id,),
            ).fetchone()[0] == 1

        invalid_delete = _request(opener, base_url + "/invoices/delete", {
            "csrf_token": "invalid",
            "id": str(invoice_id),
        })
        assert invalid_delete.status == 200
        _read(invalid_delete)
        with db_module.db() as con:
            assert con.execute("SELECT deleted_at FROM invoices WHERE id=?", (invoice_id,)).fetchone()[0] is None
        delete_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
        deleted = _request(opener, base_url + "/invoices/delete", {
            "csrf_token": _csrf(delete_html),
            "id": str(invoice_id),
        })
        assert deleted.status == 200
        _read(deleted)
        with db_module.db() as con:
            first_deleted_at = con.execute("SELECT deleted_at FROM invoices WHERE id=?", (invoice_id,)).fetchone()[0]
            assert first_deleted_at is not None
            first_audit_count = con.execute(
                "SELECT COUNT(*) FROM audit_log WHERE action='soft_delete' AND entity_type='invoices' AND entity_id=?",
                (str(invoice_id),),
            ).fetchone()[0]
            assert first_audit_count == 1
        repeated_delete_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
        repeated_delete = _request(opener, base_url + "/invoices/delete", {
            "csrf_token": _csrf(repeated_delete_html),
            "id": str(invoice_id),
        })
        assert repeated_delete.status == 200
        _read(repeated_delete)
        with db_module.db() as con:
            assert con.execute("SELECT deleted_at FROM invoices WHERE id=?", (invoice_id,)).fetchone()[0] == first_deleted_at
            assert con.execute(
                "SELECT COUNT(*) FROM audit_log WHERE action='soft_delete' AND entity_type='invoices' AND entity_id=?",
                (str(invoice_id),),
            ).fetchone()[0] == first_audit_count

        rejected = _request(opener, base_url + "/backup/create", {"csrf_token": "bad"})
        assert rejected.status == 200
        assert "CSRF rejected" not in _read(rejected).decode("utf-8")

        backup_html = _read(_request(opener, base_url + "/backup")).decode("utf-8")
        backup_token = _csrf(backup_html)
        created = _request(opener, base_url + "/backup/create", {"csrf_token": backup_token})
        assert created.status == 200
        created_html = _read(created).decode("utf-8")
        backup_name = re.search(r"/backups/([^\"']+\.sqlite3)", created_html).group(1)

        backup_content = _read(_request(opener, base_url + "/backups/" + backup_name))
        assert backup_content.startswith(b"SQLite format 3")

        restore_html = _read(_request(opener, base_url + "/backup")).decode("utf-8")
        restore_token = _csrf(restore_html)
        payload, content_type = _multipart(
            {"csrf_token": restore_token},
            {"backup_file": (backup_name, backup_content)},
        )
        restored = _request(opener, base_url + "/backup/restore", payload, {"Content-Type": content_type})
        assert restored.status == 200
        _read(restored)
        assert list(backup_dir.glob("*.sqlite3"))

        status_html = _read(_request(opener, base_url + "/status")).decode("utf-8")
        assert "Database" in status_html
        assert "OK" in status_html
        assert "Schema" in status_html
        assert "V2" in status_html
        assert "Invoice permissions" in status_html
        assert "Locking guards" in status_html
    finally:
        server.shutdown()
        server.server_close()


def test_client_contract_reference_survives_application_restart(tmp_path, monkeypatch):
    database_path = tmp_path / "data" / "client-restart.sqlite3"
    monkeypatch.setenv("PHOENIX_DB_PATH", str(database_path))
    monkeypatch.setattr(db_module, "DB_PATH", database_path)
    monkeypatch.setattr(backups, "DB_PATH", database_path)
    monkeypatch.setattr(backups, "BACKUP_DIR", tmp_path / "backups")
    db_module.invalidate_schema_cache(database_path)

    def start_server():
        server = ThreadingHTTPServer(("127.0.0.1", 0), App)
        server.daemon_threads = True
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        return server, thread, f"http://127.0.0.1:{server.server_address[1]}"

    contract_reference = "CTR persistante après redémarrage — الجزائر"
    server, thread, base_url = start_server()
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    try:
        setup_html = _read(_request(opener, base_url + "/setup")).decode("utf-8")
        _read(_request(opener, base_url + "/setup", {
            "csrf_token": _csrf(setup_html),
            "new_password": "Restart123",
            "confirm_password": "Restart123",
        }))
        client_html = _read(_request(opener, base_url + "/clients?edit_id=1")).decode("utf-8")
        _read(_request(opener, base_url + "/clients/save", {
            "csrf_token": _csrf(client_html),
            "id": "1",
            "raison_sociale": "ATM Mobilis",
            "sigle": "ATM Mobilis",
            "rgc": "RGC-M",
            "nif": "NIF-M",
            "reference_contrat": contract_reference,
        }))
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)

    server, thread, base_url = start_server()
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    try:
        login_html = _read(_request(opener, base_url + "/login")).decode("utf-8")
        _read(_request(opener, base_url + "/login", {
            "csrf_token": _csrf(login_html),
            "username": "admin",
            "password": "Restart123",
        }))
        reopened = _read(_request(opener, base_url + "/clients?edit_id=1")).decode("utf-8")
        assert contract_reference in reopened
        assert 'name="adresse"' not in reopened
        with db_module.db() as con:
            columns = {row[1] for row in con.execute("PRAGMA table_info(clients)")}
            assert "adresse" not in columns
            assert con.execute("SELECT reference_contrat FROM clients WHERE id=1").fetchone()[0] == contract_reference
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_full_invoice_workflow_runs_through_authenticated_http(tmp_path, monkeypatch):
    database_path = tmp_path / "data" / "full-workflow.sqlite3"
    uploads_path = tmp_path / "uploads"
    monkeypatch.setenv("PHOENIX_DB_PATH", str(database_path))
    monkeypatch.setattr(db_module, "DB_PATH", database_path)
    monkeypatch.setattr(backups, "DB_PATH", database_path)
    monkeypatch.setattr(backups, "BACKUP_DIR", tmp_path / "backups")
    db_module.invalidate_schema_cache(database_path)

    def isolated_upload(upload, folder, allowed_suffixes):
        suffix = app_module.Path(upload["filename"]).suffix.lower()
        if suffix not in allowed_suffixes:
            raise ValueError(f"Extension non autorisee: {suffix}")
        target_dir = uploads_path / folder
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / upload["filename"]
        target.write_bytes(upload["content"])
        return str(target)

    monkeypatch.setattr(app_module, "save_upload", isolated_upload)
    server = ThreadingHTTPServer(("127.0.0.1", 0), App)
    server.daemon_threads = True
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    opener = build_opener(HTTPCookieProcessor(CookieJar()))

    try:
        setup_html = _read(_request(opener, base_url + "/setup")).decode("utf-8")
        setup = _request(opener, base_url + "/setup", {
            "csrf_token": _csrf(setup_html),
            "new_password": "Workflow123",
            "confirm_password": "Workflow123",
        })
        assert setup.status == 200
        assert "Tableau de bord" in _read(setup).decode("utf-8")

        contract_references = [
            "CTR-2026-001",
            "CTR-2026-002-ALGER",
            "Mobilis / SAPTA / 03",
            "Référence contrat Béjaïa 04",
            "CTR—05—Tizi Ouzou",
            "عقد-قسنطينة-06",
            "Contrat École-façade-Noël-07",
            "Lot A/7; Bloc #3 (08)",
            "CTR-RN5-KM12,5-09",
            "CTR client mise à jour - الجزائر - É2E",
        ]
        for contract_reference in contract_references:
            clients_html = _read(
                _request(opener, base_url + "/clients?edit_id=1")
            ).decode("utf-8")
            assert 'name="adresse"' not in clients_html
            assert 'name="reference_contrat"' in clients_html
            client_settings = _request(
                opener,
                base_url + "/clients/save",
                {
                    "csrf_token": _csrf(clients_html),
                    "id": "1",
                    "raison_sociale": "ATM Mobilis",
                    "sigle": "ATM Mobilis",
                    "rgc": "RGC-M",
                    "nif": "000099999999999",
                    "reference_contrat": contract_reference,
                },
            )
            assert client_settings.status == 200
            assert "ATM Mobilis" in _read(client_settings).decode("utf-8")
            with db_module.db() as con:
                assert con.execute(
                    "SELECT reference_contrat FROM clients WHERE id=1"
                ).fetchone()[0] == contract_reference

        direction_html = _read(
            _request(opener, base_url + "/clients?client_id=1&new_direction=1")
        ).decode("utf-8")
        direction_response = _request(
            opener,
            base_url + "/clients/directions/save",
            {
                "csrf_token": _csrf(direction_html),
                "client_id": "1",
                "company_branch_id": "1",
                "sigle": "DR E2E",
                "name": "Direction E2E",
                "address": "Adresse direction mise à jour - الجزائر - É2E",
            },
        )
        assert direction_response.status == 200
        assert "Direction E2E" in _read(direction_response).decode("utf-8")
        with db_module.db() as con:
            direction_id = con.execute(
                "SELECT id FROM client_directions WHERE name='Direction E2E'"
            ).fetchone()[0]
            client_sigle = con.execute("SELECT sigle FROM clients WHERE id=1").fetchone()[0]
            direction_sigle = con.execute("SELECT sigle FROM client_directions WHERE id=?", (direction_id,)).fetchone()[0]
            branch_sigle = con.execute("SELECT sigle FROM company_branches WHERE id=1").fetchone()[0]

        partner_workbook = load_workbook(BytesIO(partner_template("subcontractors")))
        partner_workbook.active.append(["ST Import E2E", "Contact E2E", "0550000000", "contact@e2e.dz", "Alger", "Actif"])
        partner_stream = BytesIO()
        partner_workbook.save(partner_stream)
        partner_html = _read(_request(opener, base_url + "/settings?section=subcontractors")).decode("utf-8")
        partner_payload, partner_content_type = _multipart(
            {"csrf_token": _csrf(partner_html), "kind": "subcontractors"},
            {"excel_file": ("sous-traitants.xlsx", partner_stream.getvalue())},
        )
        partner_preview_response = _request(opener, base_url + "/settings/partners/import/preview", partner_payload, {"Content-Type": partner_content_type})
        partner_preview_html = _read(partner_preview_response).decode("utf-8")
        partner_token = parse_qs(urlparse(partner_preview_response.geturl()).query)["import_token"][0]
        assert "ST Import E2E" in partner_preview_html and "Confirmer l’import" in partner_preview_html
        partner_confirm = _request(opener, base_url + "/settings/partners/import/confirm", {
            "csrf_token": _csrf(partner_preview_html), "kind": "subcontractors", "token": partner_token,
        })
        assert "ST Import E2E" in _read(partner_confirm).decode("utf-8")

        po_workbook = load_workbook(BytesIO(purchase_order_template()))
        po_workbook["Bons_de_commande"].append([
            "BC-EXCEL-E2E", "10/08/2026", client_sigle, direction_sigle,
            branch_sigle, "CONST", 1190, "Import E2E",
        ])
        po_workbook["Sites"].append([
            "BC-EXCEL-E2E", "SITE-EXCEL-E2E", "Site Excel E2E", "A12",
            "A12 ( MAT 12M + BTS OUTDOOR )", "ST Import E2E", "",
        ])
        po_stream = BytesIO()
        po_workbook.save(po_stream)
        po_page = _read(_request(opener, base_url + "/purchase-orders")).decode("utf-8")
        po_import_payload, po_import_type = _multipart(
            {"csrf_token": _csrf(po_page)},
            {"excel_file": ("bc-sites.xlsx", po_stream.getvalue())},
        )
        po_preview_response = _request(opener, base_url + "/purchase-orders/import/preview", po_import_payload, {"Content-Type": po_import_type})
        po_preview_html = _read(po_preview_response).decode("utf-8")
        po_token = parse_qs(urlparse(po_preview_response.geturl()).query)["import_token"][0]
        assert "BC-EXCEL-E2E" in po_preview_html and "Confirmer l’import" in po_preview_html
        po_confirm = _request(opener, base_url + "/purchase-orders/import/confirm", {
            "csrf_token": _csrf(po_preview_html), "token": po_token,
        })
        assert "BC-EXCEL-E2E" in _read(po_confirm).decode("utf-8")
        exported_bc = _request(opener, base_url + "/purchase-orders/export")
        exported_book = load_workbook(BytesIO(_read(exported_bc)), data_only=True)
        assert exported_book.sheetnames == ["Bons_de_commande", "Sites"]

        bpu_workbook = make_xlsx([
            ("BPU", [
                ["Item", "Designation", "Unite", "Prix"],
                ["", "Fourniture", "", ""],
                [1, "Cable E2E", "m", 100],
                ["", "Prestation", "", ""],
                [2, "Pose E2E", "U", 200],
            ])
        ])
        bpu_html = _read(_request(opener, base_url + "/bpu")).decode("utf-8")
        bpu_payload, bpu_content_type = _multipart(
            {"csrf_token": _csrf(bpu_html)},
            {"bpu_file": ("bpu-e2e.xlsx", bpu_workbook)},
        )
        bpu_response = _request(
            opener,
            base_url + "/bpu",
            bpu_payload,
            {"Content-Type": bpu_content_type},
        )
        assert bpu_response.status == 200
        bpu_response_html = _read(bpu_response).decode("utf-8")
        assert "2 articles BPU importes" in bpu_response_html
        assert "Cable E2E" in bpu_response_html

        purchase_order_html = _read(
            _request(opener, base_url + "/purchase-orders?new=1")
        ).decode("utf-8")
        po_payload, po_content_type = _multipart(
            {
                "csrf_token": _csrf(purchase_order_html),
                "numero_bc": "BC-MGC-E2E-FULL",
                "date_bc": "2026-08-01",
                "client_direction_id": str(direction_id),
                "type_bc": "MGC",
                "objet": "Maintenance genie civil E2E",
                "montant_ttc": "2261",
                "code_sites": "",
            },
            {},
        )
        po_response = _request(
            opener,
            base_url + "/purchase-orders",
            po_payload,
            {"Content-Type": po_content_type},
        )
        assert po_response.status == 200
        assert "BC-MGC-E2E-FULL" in _read(po_response).decode("utf-8")
        with db_module.db() as con:
            po = con.execute(
                "SELECT id, type_bc FROM purchase_orders WHERE numero_bc='BC-MGC-E2E-FULL'"
            ).fetchone()
            assert po["type_bc"] == "MGC"
            po_id = po["id"]
            subcontractor_id = con.execute(
                "INSERT INTO subcontractors(raison_sociale,sigle) VALUES('ST Workflow','ST WORKFLOW')"
            ).lastrowid

        po_detail_html = _read(
            _request(opener, base_url + f"/purchase-orders?view_id={po_id}")
        ).decode("utf-8")
        site_response = _request(opener, base_url + "/sites", {
            "csrf_token": _csrf(po_detail_html),
            "purchase_order_id": str(po_id),
            "code_site": "SITE-MGC-E2E-FULL",
            "nom_site": "Site MGC E2E complet",
            "typologie_site": "MGC",
            "typologie_libelle": "MGC",
            "subcontractor_id": str(subcontractor_id),
            "design_office_id": "",
        })
        assert site_response.status == 200
        assert "SITE-MGC-E2E-FULL" in _read(site_response).decode("utf-8")
        with db_module.db() as con:
            site_id = con.execute(
                "SELECT id FROM sites WHERE code_site='SITE-MGC-E2E-FULL'"
            ).fetchone()[0]

        invoice_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
        assert "BPU ENT" in invoice_html
        assert "BPU ST" in invoice_html
        assert "Numérotation des articles" not in invoice_html
        assert invoice_html.count("data-searchable-select") == 1
        assert 'id="invoice-site-display" readonly' in invoice_html
        assert 'id="invoice-site-select" data-placeholder="Saisir ou rechercher un code site"' in invoice_html
        assert invoice_html.index('<span>Bon de commande</span>') < invoice_html.index('<span>Site</span>')
        assert invoice_html.index('<span>Site</span>') < invoice_html.index('<span>N° Facture</span>')
        assert invoice_html.index('<span>N° Facture</span>') < invoice_html.index('<span>Date facture</span>')
        assert invoice_html.index('<span>Date facture</span>') < invoice_html.index('<span>Nature facture</span>')
        assert invoice_html.index('<span>Nature facture</span>') < invoice_html.index('<span>Typologie</span>')
        assert 'id="invoice-lines-scroll"' in invoice_html
        assert invoice_html.count('class="calculation-stage"') == 1
        assert invoice_html.count('calculation-stage') >= 2
        assert "function ensureTrailingBlankLine()" in invoice_html
        assert "completedCount >= 10" in invoice_html
        assert "addInvoiceLine(true);" in invoice_html
        assert "for (let index = 0; index < 4" not in invoice_html

        ui_script = _read(_request(opener, base_url + "/static/app-ui.js")).decode("utf-8")
        assert "success: 5000" in ui_script
        assert "warning: 7000" in ui_script
        assert "error: 10000" in ui_script
        assert "mouseenter" in ui_script and "mouseleave" in ui_script
        invoice_response = _request(opener, base_url + "/invoices", {
            "csrf_token": _csrf(invoice_html),
            "invoice_number": "FAC-E2E-FULL",
            "invoice_date": "2026-08-02",
            "purchase_order_id": str(po_id),
            "site_id": str(site_id),
            "lines": json.dumps([
                {"article_number": 1, "quantity": 10, "source_reference_type": "ST", "source_st_number": 1},
                {"article_number": 2, "quantity": 5, "source_reference_type": "GENERAL", "source_st_number": None},
            ]),
        })
        assert invoice_response.status == 200
        invoice_response_html = _read(invoice_response).decode("utf-8")
        assert "Prête pour dépôt DTC" in invoice_response_html
        assert "CONST" in invoice_response_html
        with db_module.db() as con:
            invoice = con.execute(
                """
                SELECT id, invoice_type, total_ht, total_ttc, lifecycle_status, numbering_system
                FROM invoice_lifecycle WHERE invoice_number='FAC-E2E-FULL'
                """
            ).fetchone()
            assert tuple(invoice)[1:] == ("CONSTRUCTION", 2000, 2261, "READY_DTC", "ST")
            invoice_id = invoice["id"]
            assert [tuple(row) for row in con.execute(
                "SELECT article_number, source_reference_type, source_st_number FROM invoice_lines WHERE invoice_id=? ORDER BY id",
                (invoice_id,),
            )] == [(1, "ST", 1), (2, "GENERAL", None)]

        tracking_html = _read(
            _request(opener, base_url + "/table-facturation-new")
        ).decode("utf-8")
        tracking_token = _csrf(tracking_html)
        for field, value, expected in (
            ("date_depot_dtc", "2026-08-03", "DEPOSITED_DTC"),
            ("date_depot_mobilis", "2026-08-04", "AWAITING_PAYMENT"),
            ("date_ov", "2026-08-05", "PAID"),
        ):
            response = _request(
                opener,
                base_url + "/table-facturation-new/update",
                {
                    "csrf_token": tracking_token,
                    "invoice_id": str(invoice_id),
                    field: value,
                },
            )
            assert response.status == 200
            assert expected in _read(response).decode("utf-8")

        payment_reference = _request(
            opener,
            base_url + "/table-facturation-new/update",
            {
                "csrf_token": tracking_token,
                "invoice_id": str(invoice_id),
                "numero_ordre_virement": "OV-FULL-2026-77",
            },
        )
        assert payment_reference.status == 200
        assert "OV-FULL-2026-77" in _read(payment_reference).decode("utf-8")

        paid_table_html = _read(
            _request(opener, base_url + "/table-facturation-new?status=PAID")
        ).decode("utf-8")
        assert "FAC-E2E-FULL" in paid_table_html
        assert "Payée" in paid_table_html
        assert "OV-FULL-2026-77" in paid_table_html

        table_export = _request(
            opener,
            base_url + "/table-facturation-new/export?status=PAID",
        )
        assert table_export.status == 200
        table_book = load_workbook(BytesIO(_read(table_export)), data_only=True, read_only=True)
        try:
            table_sheet = table_book["Table Facturation"]
            headers = [cell.value for cell in next(table_sheet.iter_rows())]
            assert headers == [
                "DR MOBILIS", "N° BC", "CODE SITE", "TYPE BC", "OBJET BC",
                "BET", "ST", "TYPOLOGIE SITE", "N° FACTURE", "DATE FACTURE",
                "TTC FACTURE", "DATE DÉPÔT DTC", "DATE DÉPÔT MOBILIS",
                "DATE PAIEMENT", "N° ORDRE DE VIREMENT", "ÉTAT", "REMARQUE",
            ]
            values = [cell.value for cell in next(table_sheet.iter_rows(min_row=2))]
            assert "OV-FULL-2026-77" in values
        finally:
            table_book.close()

        dashboard_html = _read(_request(opener, base_url + "/")).decode("utf-8")
        assert "Payée" in dashboard_html
        assert "2,26 KDA" in dashboard_html

        excel = _request(opener, base_url + f"/invoices/export?id={invoice_id}")
        assert excel.status == 200
        assert excel.headers["Content-Type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        excel_content = _read(excel)
        assert excel_content.startswith(b"PK")
        workbook = load_workbook(BytesIO(excel_content), data_only=True, read_only=True)
        try:
            assert any(
                cell.value == "Adresse direction mise à jour - الجزائر - É2E"
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
            )
        finally:
            workbook.close()

        preview = _request(
            opener,
            base_url + f"/invoices/preview/facture?id={invoice_id}",
        )
        assert preview.status == 200
        preview_html = _read(preview).decode("utf-8")
        assert "FAC-E2E-FULL" in preview_html
        assert "Adresse direction mise à jour - الجزائر - É2E" in preview_html
        assert "CTR client mise à jour - الجزائر - É2E" in preview_html
        pdf = _request(
            opener,
            base_url + f"/invoices/pdf/facture/{invoice_id}/facture.pdf",
            timeout=60,
        )
        assert pdf.status == 200
        assert pdf.headers["Content-Type"] == "application/pdf"
        pdf_content = _read(pdf)
        assert pdf_content.startswith(b"%PDF")
        assert len(pdf_content) > 10_000

        with db_module.db() as con:
            tracking = con.execute(
                """
                SELECT date_depot_dtc, date_depot_mobilis, date_ov, numero_ordre_virement
                FROM invoice_tracking WHERE invoice_id=?
                """,
                (invoice_id,),
            ).fetchone()
            assert tuple(tracking) == (
                "2026-08-03", "2026-08-04", "2026-08-05", "OV-FULL-2026-77"
            )
            assert con.execute(
                """
                SELECT COUNT(*) FROM invoice_tracking_events
                WHERE invoice_id=? AND event_type='INVOICE_EXPORTED'
                  AND new_value IN ('invoice_xlsx', 'facture_pdf')
                """,
                (invoice_id,),
            ).fetchone()[0] == 2
            assert con.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert con.execute("PRAGMA foreign_key_check").fetchall() == []
    finally:
        server.shutdown()
        server.server_close()
