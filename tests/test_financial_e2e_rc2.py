from decimal import Decimal, ROUND_DOWN, localcontext
from http.cookiejar import CookieJar
from http.server import ThreadingHTTPServer
from io import BytesIO
import json
import re
import threading
from zipfile import ZipFile
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, Request, build_opener

from openpyxl import load_workbook
import pytest

import db as db_module
from app import App
from services import backups


pytestmark = pytest.mark.e2e
Q = Decimal("0.01")


def _truncate(value):
    with localcontext() as context:
        context.prec = 50
        return Decimal(value).quantize(Q, rounding=ROUND_DOWN)


def _oracle(quantity, price, rg_rate, tva_rate):
    line = _truncate(Decimal(quantity) * Decimal(price))
    ht = _truncate(line)
    rg = _truncate(ht * Decimal(rg_rate) / Decimal("100"))
    after = _truncate(ht - rg)
    tva = _truncate(after * Decimal(tva_rate) / Decimal("100"))
    return line, (ht, rg, after, tva, _truncate(after + tva))


def _csrf(html):
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    if not match:
        match = re.search(r'name="csrf-token" content="([^"]+)"', html)
    assert match
    return match.group(1)


def _read(response):
    try:
        return response.read()
    finally:
        response.close()


def _request(opener, url, data=None, timeout=60):
    payload = urlencode(data).encode("utf-8") if isinstance(data, dict) else data
    return opener.open(Request(url, data=payload), timeout=timeout)


def _start_server():
    server = ThreadingHTTPServer(("127.0.0.1", 0), App)
    server.daemon_threads = True
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    return server, thread, f"http://127.0.0.1:{server.server_address[1]}"


def _stop_server(server, thread):
    server.shutdown()
    server.server_close()
    thread.join(timeout=10)


def _set_global_rates(connection, rg_fraction, tva_fraction):
    for key, value in (("retention_rate", rg_fraction), ("tax_rate", tva_fraction)):
        connection.execute(
            """
            INSERT INTO app_settings(key, value, updated_by) VALUES(?, ?, 'financial-e2e')
            ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_by=excluded.updated_by
            """,
            (key, value),
        )


def _xlsx_facture(content):
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        rows = [tuple(cell.value for cell in row) for row in workbook["Facture"].iter_rows()]
        return [value for row in rows for value in row], rows
    finally:
        workbook.close()


def _money_fr(value):
    whole, cents = f"{value:.2f}".split(".")
    groups = []
    while whole:
        groups.append(whole[-3:])
        whole = whole[:-3]
    return f"{' '.join(reversed(groups))},{cents}"


def test_ten_invoice_snapshot_restart_pdf_xlsx_reconciliation(tmp_path, monkeypatch):
    database_path = tmp_path / "financial-e2e" / "pos_ai.sqlite3"
    monkeypatch.setenv("PHOENIX_DB_PATH", str(database_path))
    monkeypatch.setattr(db_module, "DB_PATH", database_path)
    monkeypatch.setattr(backups, "DB_PATH", database_path)
    monkeypatch.setattr(backups, "BACKUP_DIR", tmp_path / "backups")
    db_module.invalidate_schema_cache(database_path)

    server, thread, base_url = _start_server()
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    invoice_cases = []
    try:
        setup_html = _read(_request(opener, base_url + "/setup")).decode("utf-8")
        _read(_request(opener, base_url + "/setup", {
            "csrf_token": _csrf(setup_html),
            "new_password": "Finance123",
            "confirm_password": "Finance123",
        }))

        with db_module.db() as connection:
            direction_id = connection.execute(
                "SELECT id FROM client_directions ORDER BY id LIMIT 1"
            ).fetchone()[0]
            branch_id = connection.execute(
                "SELECT id FROM company_branches ORDER BY id LIMIT 1"
            ).fetchone()[0]
            connection.execute(
                """
                INSERT OR REPLACE INTO bpu_items(
                    article_number, designation, unite, pu_ht, categorie
                ) VALUES(9901, 'Article financier RC2', 'U', 12.349, 'acquisition')
                """
            )
            design_office_id = connection.execute(
                "INSERT INTO design_offices(raison_sociale, sigle) VALUES('BET Finance RC2', 'BET FIN RC2')"
            ).lastrowid
            _set_global_rates(connection, "0.05", "0.19")
            for index in range(10):
                po_id = connection.execute(
                    """
                    INSERT INTO purchase_orders(
                        numero_bc, date_bc, client_direction_id, company_branch_id,
                        type_bc, objet, montant_ttc, created_by, updated_by
                    ) VALUES(?, '2026-08-25', ?, ?, 'ACQUISITION', 'RC2 financial E2E',
                             1000000, 'financial-e2e', 'financial-e2e')
                    """,
                    (f"QA-RC2-BC-{index + 1:02d}", direction_id, branch_id),
                ).lastrowid
                site_id = connection.execute(
                    """
                    INSERT INTO sites(
                        purchase_order_id, code_site, nom_site, typologie_site,
                        typology_id, design_office_id, created_by, updated_by
                    ) VALUES(?, ?, ?, 'A12',
                        (SELECT id FROM typologies WHERE sigle='A12'), ?,
                        'financial-e2e', 'financial-e2e')
                    """,
                    (po_id, f"QA-RC2-SITE-{index + 1:02d}", f"Site financier {index + 1}", design_office_id),
                ).lastrowid
                invoice_cases.append({"index": index, "po_id": po_id, "site_id": site_id})

        for case in invoice_cases:
            if case["index"] == 5:
                with db_module.db() as connection:
                    _set_global_rates(connection, "0.10", "0.20")
            rg_rate, tva_rate = (("5", "19") if case["index"] < 5 else ("10", "20"))
            initial_quantity = f"{case['index'] + 1}.125"
            invoice_html = _read(_request(opener, base_url + "/invoices")).decode("utf-8")
            response_html = _read(_request(opener, base_url + "/invoices", {
                "csrf_token": _csrf(invoice_html),
                "invoice_number": f"QA-RC2-FAC-{case['index'] + 1:02d}",
                "invoice_date": f"2026-08-{case['index'] + 1:02d}",
                "purchase_order_id": str(case["po_id"]),
                "site_id": str(case["site_id"]),
                "lines": json.dumps([{"article_number": 9901, "quantity": initial_quantity}]),
            })).decode("utf-8")
            assert "Enregistrement" in response_html, re.findall(
                r'<section class="alert"[^>]*>(.*?)</section>', response_html, re.DOTALL
            )
            with db_module.db() as connection:
                invoice = connection.execute(
                    "SELECT id, rg_rate, tva_rate FROM invoices WHERE invoice_number=?",
                    (f"QA-RC2-FAC-{case['index'] + 1:02d}",),
                ).fetchone()
                assert tuple(invoice[1:]) == (int(rg_rate), int(tva_rate))
                case.update(invoice_id=invoice["id"], rg_rate=rg_rate, tva_rate=tva_rate)

            final_quantity = f"{case['index'] + 1}.349"
            edit_html = _read(
                _request(opener, base_url + f"/invoices?edit_id={case['invoice_id']}")
            ).decode("utf-8")
            assert f"RG {rg_rate},00%" in edit_html
            assert f"TVA {tva_rate},00%" in edit_html
            _read(_request(opener, base_url + "/invoices", {
                "csrf_token": _csrf(edit_html),
                "id": str(case["invoice_id"]),
                "invoice_number": f"QA-RC2-FAC-{case['index'] + 1:02d}",
                "invoice_date": f"2026-08-{case['index'] + 1:02d}",
                "purchase_order_id": str(case["po_id"]),
                "site_id": str(case["site_id"]),
                "lines": json.dumps([{"article_number": 9901, "quantity": final_quantity}]),
            }))
            case["quantity"] = final_quantity

        with db_module.db() as connection:
            _set_global_rates(connection, "0.07", "0.21")
            for case in invoice_cases:
                invoice = connection.execute(
                    """
                    SELECT rg_rate, tva_rate, total_ht, retenue_garantie,
                           montant_ht_apres_rg, tva, total_ttc
                    FROM invoices WHERE id=?
                    """,
                    (case["invoice_id"],),
                ).fetchone()
                _, expected = _oracle(case["quantity"], "12.349", case["rg_rate"], case["tva_rate"])
                assert Decimal(str(invoice["rg_rate"])) == Decimal(case["rg_rate"])
                assert Decimal(str(invoice["tva_rate"])) == Decimal(case["tva_rate"])
                assert tuple(Decimal(str(value)) for value in invoice[2:]) == expected
                case["expected"] = expected

        _stop_server(server, thread)
        server = thread = None
        server, thread, base_url = _start_server()
        opener = build_opener(HTTPCookieProcessor(CookieJar()))
        login_html = _read(_request(opener, base_url + "/login")).decode("utf-8")
        _read(_request(opener, base_url + "/login", {
            "csrf_token": _csrf(login_html), "username": "admin", "password": "Finance123",
        }))

        for case in invoice_cases:
            edit_html = _read(
                _request(opener, base_url + f"/invoices?edit_id={case['invoice_id']}")
            ).decode("utf-8")
            assert f"RG {case['rg_rate']},00%" in edit_html
            assert f"TVA {case['tva_rate']},00%" in edit_html

            preview_html = _read(
                _request(opener, base_url + f"/invoices/preview/facture?id={case['invoice_id']}")
            ).decode("utf-8")
            assert f"RETENUE DE GARANTIE {case['rg_rate']}%" in preview_html
            assert f"T V A {case['tva_rate']}%" in preview_html
            for value in case["expected"]:
                assert _money_fr(value) in preview_html

            xlsx_content = _read(
                _request(opener, base_url + f"/invoices/export?id={case['invoice_id']}")
            )
            xlsx_dir = tmp_path / "generated-xlsx"
            xlsx_dir.mkdir(exist_ok=True)
            (xlsx_dir / f"invoice-{case['index'] + 1:02d}.xlsx").write_bytes(xlsx_content)
            with ZipFile(BytesIO(xlsx_content)) as archive:
                facture_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
            assert "<v>12.34</v>" in facture_xml
            assert "00000000000001" not in facture_xml
            for expected_value in case["expected"]:
                assert f"<v>{expected_value:.2f}</v>" in facture_xml
            values, rows = _xlsx_facture(xlsx_content)
            assert f"RETENUE DE GARANTIE {case['rg_rate']}%" in values, [
                value for value in values if isinstance(value, str) and "GARANTIE" in value
            ]
            assert f"TVA {case['tva_rate']} %" in values, [
                value for value in values if isinstance(value, str) and "TVA" in value
            ]
            financial_rows = {
                re.sub(r"[^A-Z0-9]", "", row[4].upper()): row[5]
                for row in rows
                if len(row) >= 6 and isinstance(row[4], str)
                and any(term in row[4] for term in ("TOTAL EN HT", "RETENUE", "MONTANT HT", "TVA", "T T C", "H.T"))
            }
            expected_labels = (
                "TOTALENHT",
                f"RETENUEDEGARANTIE{case['rg_rate']}",
                "MONTANTHTAPRESRG",
                f"TVA{case['tva_rate']}",
                "TOTALENTTC",
            )
            for label, expected_value in zip(expected_labels, case["expected"]):
                assert Decimal(str(financial_rows[label])) == expected_value

            pdf_content = _read(
                _request(
                    opener,
                    base_url + f"/invoices/pdf/facture/{case['invoice_id']}/facture.pdf",
                    timeout=90,
                )
            )
            assert pdf_content.startswith(b"%PDF") and len(pdf_content) > 10_000
            pdf_dir = tmp_path / "generated-pdfs"
            pdf_dir.mkdir(exist_ok=True)
            (pdf_dir / f"invoice-{case['index'] + 1:02d}.pdf").write_bytes(pdf_content)

        with db_module.db() as connection:
            assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
            assert connection.execute(
                "SELECT COUNT(*) FROM invoices WHERE rg_rate IS NULL OR tva_rate IS NULL"
            ).fetchone()[0] == 0
            assert connection.execute(
                "SELECT COUNT(*) FROM invoices WHERE invoice_number LIKE 'QA-RC2-FAC-%'"
            ).fetchone()[0] == 10
    finally:
        if server is not None:
            _stop_server(server, thread)
        db_module.invalidate_schema_cache(database_path)
